import json
import os
import random
import re
import sys
import threading
import time
import traceback
import urllib.error
import urllib.request
import uuid
import asyncio
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).parent.resolve()
STATIC_DIR = ROOT / "static"
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(exist_ok=True)
CURRENT_BATCH_PATH = DATA_DIR / "current_batch.json"
CURRENT_JOB_PATH = DATA_DIR / "current_job.json"
SERVICE_CONFIG_PATH = DATA_DIR / "service_config.json"

STYLE_NAMES = [
    "americana",
    "athleisure",
    "avant_garde",
    "bohemian",
    "casual",
    "classic",
    "clean",
    "coquette",
    "cottagecore",
    "cozy",
    "dark_academia",
    "edgy",
    "elegant",
    "gorpcore",
    "luxe",
    "minimal",
    "moto",
    "premium",
    "preppy",
    "quiet_luxury",
    "relaxed",
    "resort",
    "romantic",
    "sport",
    "streetwear",
    "sustainable",
    "tailored",
    "technical",
    "vintage",
    "western",
    "workwear",
    "y2k",
]

EXCEL_PATH = Path(
    os.getenv(
        "VIDEO_EXCEL_PATH",
        "/Users/guyin/Downloads/video_urls_2026-05-24 (1)_with_caption_hashtag.xlsx",
    )
)
TEXT_API_URL = os.getenv(
    "EVOLINK_TEXT_API_URL",
    "https://direct.evolink.ai/v1/chat/completions",
)
TEXT_MODEL = os.getenv("EVOLINK_TEXT_MODEL", "gemini-3.5-flash")
VIDEO_API_URL = os.getenv(
    "EVOLINK_VIDEO_API_URL",
    "https://direct.evolink.ai/v1/chat/completions",
)
def read_env_file_value(text, key):
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        name, value = line.split("=", 1)
        if name.strip() != key:
            continue
        value = value.strip()
        if (value.startswith('"') and value.endswith('"')) or (
            value.startswith("'") and value.endswith("'")
        ):
            value = value[1:-1]
        return value
    return ""


def load_secret(key):
    value = os.getenv(key, "")
    if value:
        return value
    env_path = ROOT / ".env"
    if env_path.exists():
        return read_env_file_value(env_path.read_text(errors="ignore"), key)
    return ""


API_KEY = load_secret("EVOLINK_API_KEY")
DATABASE_URL = load_secret("DATABASE_URL")
DEFAULT_API_CONCURRENCY = 200
DEFAULT_BLOGGER_MIN_VIDEO_COUNT = 15
DEFAULT_VIDEO_WORKER_COUNT = 20
DEFAULT_BLOGGER_WORKER_COUNT = 5
DEFAULT_QUEUE_POLL_SECONDS = 2
DEFAULT_TASK_LOCK_SECONDS = 3600
GCS_SIGNED_URL_TTL = timedelta(days=7)
GCS_SIGNED_URL_REFRESH_THRESHOLD = timedelta(days=1)
GCS_HTTPS_HOSTS = {"storage.googleapis.com", "storage.cloud.google.com"}

JOBS = {}
JOBS_LOCK = threading.Lock()
BATCH_LOCK = threading.Lock()
QUEUE_WORKERS_LOCK = threading.Lock()
QUEUE_WORKERS_STARTED = False


PROMPT_DESCRIPTIONS = {
    "1": "负责生成单视频描述单元：从视频、caption、hashtag 提取可用于后续账号判断的客观证据。",
    "2": "负责生成账号基础人口、消费层级、气质心理：只吃 15 个 video_description_unit，不直接生成 social_identity / occasion / style。",
    "3": "负责单视频 10 属性分类：基础人口、消费层级、气质心理、社会身份、Occasion 等单视频标签。",
    "4": "负责输出 32 维风格向量：严格按 STYLE_DIMENSIONS 输出 0-1 分数。",
    "5": "负责输出单视频 StyleSignature：8-facet 细粒度美学指纹。",
    "6": "负责生成博主一句话总结：基于 TikTok profile/bio 和所有 video_description_unit 输出 account_one_sentence_summary。",
}


def read_service_config():
    if not SERVICE_CONFIG_PATH.exists():
        return {}
    try:
        data = json.loads(SERVICE_CONFIG_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def write_service_config(config):
    SERVICE_CONFIG_PATH.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")


def mask_secret(value):
    if not value:
        return ""
    if len(value) <= 10:
        return "***"
    return f"{value[:5]}...{value[-4:]}"


def service_config_value(key, default=""):
    return read_service_config().get(key) or default


def current_api_key():
    return service_config_value("api_key", API_KEY)


def current_text_api_url():
    return service_config_value("text_api_url", TEXT_API_URL)


def current_text_model():
    return service_config_value("text_model", TEXT_MODEL)


def current_video_api_url():
    return service_config_value("video_api_url", VIDEO_API_URL)


def current_api_concurrency():
    try:
        return max(1, min(int(service_config_value("default_api_concurrency", DEFAULT_API_CONCURRENCY)), 500))
    except Exception:
        return DEFAULT_API_CONCURRENCY


def current_blogger_min_video_count():
    try:
        return max(1, min(int(service_config_value("blogger_min_video_count", DEFAULT_BLOGGER_MIN_VIDEO_COUNT)), 50))
    except Exception:
        return DEFAULT_BLOGGER_MIN_VIDEO_COUNT


def current_video_worker_count():
    try:
        value = service_config_value("video_worker_count", os.getenv("VIDEO_WORKER_COUNT", DEFAULT_VIDEO_WORKER_COUNT))
        return max(1, min(int(value), 100))
    except Exception:
        return DEFAULT_VIDEO_WORKER_COUNT


def current_blogger_worker_count():
    try:
        value = service_config_value("blogger_worker_count", os.getenv("BLOGGER_WORKER_COUNT", DEFAULT_BLOGGER_WORKER_COUNT))
        return max(1, min(int(value), 50))
    except Exception:
        return DEFAULT_BLOGGER_WORKER_COUNT


def current_queue_poll_seconds():
    try:
        return max(1, min(int(os.getenv("QUEUE_POLL_SECONDS", DEFAULT_QUEUE_POLL_SECONDS)), 30))
    except Exception:
        return DEFAULT_QUEUE_POLL_SECONDS


def public_config_payload():
    config = read_service_config()
    effective_key = current_api_key()
    prompts = {
        str(number): load_default_prompt(number)
        for number in range(1, 7)
    }
    return {
        "text_api_url": current_text_api_url(),
        "text_model": current_text_model(),
        "video_api_url": current_video_api_url(),
        "video_model": "gemini-3.1-pro-preview",
        "has_api_key": bool(effective_key),
        "api_key_masked": mask_secret(effective_key),
        "default_api_concurrency": current_api_concurrency(),
        "blogger_min_video_count": current_blogger_min_video_count(),
        "video_worker_count": current_video_worker_count(),
        "blogger_worker_count": current_blogger_worker_count(),
        "prompts": prompts,
        "prompt_descriptions": PROMPT_DESCRIPTIONS,
        "config_source": "service_config" if config else "defaults",
    }


def update_service_config(body):
    config = read_service_config()
    for key in ("text_api_url", "text_model", "video_api_url"):
        if key in body:
            config[key] = str(body.get(key) or "").strip()
    if str(body.get("api_key") or "").strip():
        config["api_key"] = str(body["api_key"]).strip()
    if "default_api_concurrency" in body:
        try:
            config["default_api_concurrency"] = max(1, min(int(body["default_api_concurrency"]), 500))
        except Exception:
            config["default_api_concurrency"] = current_api_concurrency()
    if "blogger_min_video_count" in body:
        try:
            config["blogger_min_video_count"] = max(1, min(int(body["blogger_min_video_count"]), 50))
        except Exception:
            config["blogger_min_video_count"] = current_blogger_min_video_count()
    if "video_worker_count" in body:
        try:
            config["video_worker_count"] = max(1, min(int(body["video_worker_count"]), 100))
        except Exception:
            config["video_worker_count"] = current_video_worker_count()
    if "blogger_worker_count" in body:
        try:
            config["blogger_worker_count"] = max(1, min(int(body["blogger_worker_count"]), 50))
        except Exception:
            config["blogger_worker_count"] = current_blogger_worker_count()
    if isinstance(body.get("prompts"), dict):
        prompts = config.get("prompts") if isinstance(config.get("prompts"), dict) else {}
        for key in ("1", "2", "3", "4", "5", "6"):
            if key in body["prompts"]:
                prompts[key] = str(body["prompts"][key] or "")
        config["prompts"] = prompts
    write_service_config(config)
    return public_config_payload()


def read_json_body(handler):
    length = int(handler.headers.get("Content-Length", "0") or "0")
    if length == 0:
        return {}
    return json.loads(handler.rfile.read(length).decode("utf-8"))


def json_response(handler, payload, status=200):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def error_response(handler, message, status=400):
    json_response(handler, {"error": message}, status)


def api_code_response(handler, code, message, http_status=200, **extra):
    json_response(handler, {"code": code, "message": message, **extra}, http_status)


def beijing_day_to_utc_range(day):
    if not day:
        return None, None
    try:
        local_start = datetime.strptime(day, "%Y-%m-%d").replace(tzinfo=timezone(timedelta(hours=8)))
    except Exception:
        raise ValueError("date must be YYYY-MM-DD")
    return local_start.astimezone(timezone.utc), (local_start + timedelta(days=1)).astimezone(timezone.utc)


def validate_video_tagging_request(body):
    disallowed = [key for key in ("source_url", "request_id", "local_video_url") if key in body]
    if disallowed:
        return False, {"code": 4006, "message": f"unsupported fields: {', '.join(disallowed)}"}
    required = [
        ("video_id", 4001),
        ("gcs_url", 4002),
        ("description", 4003),
        ("callback_url", 4004),
    ]
    for key, code in required:
        if not str(body.get(key) or "").strip():
            return False, {"code": code, "message": f"{key} is required"}
    try:
        uuid.UUID(str(body["video_id"]))
    except Exception:
        return False, {"code": 4005, "message": "video_id must be a UUID"}
    callback_url = str(body["callback_url"]).strip()
    if not callback_url.startswith(("http://", "https://")):
        return False, {"code": 4004, "message": "callback_url must start with http:// or https://"}
    return True, None


def video_tagging_table_sql():
    return """
create table if not exists public.video_tagging_results (
  id uuid primary key,
  video_id uuid not null unique,
  gcs_url text not null,
  description text not null,
  status varchar not null,
  result_code integer,
  result_message text,
  error_code varchar,
  error_message text,
  video_description_unit jsonb,
  personal_tags jsonb,
  style_vector jsonb,
  style_signature jsonb,
  raw_outputs jsonb,
  callback_url text not null,
  callback_status varchar,
  callback_response_code integer,
  callback_response_body text,
  callback_attempts integer not null default 0,
  source_type varchar not null default 'direct',
  source_blogger_task_id uuid,
  source_tiktok_blogger_id uuid,
  last_callback_at timestamptz,
  created_at timestamptz not null default now(),
  updated_at timestamptz not null default now(),
  started_at timestamptz,
  finished_at timestamptz
);
"""


def build_video_tagging_callback_payload(task):
    status = task.get("status")
    is_success = status == "success"
    return {
        "event": "video_tagging.completed" if is_success else "video_tagging.failed",
        "video_id": str(task.get("video_id")),
        "task_id": str(task.get("id")),
        "status": status,
        "code": int(task.get("code") if task.get("code") is not None else task.get("result_code") or (0 if is_success else 5000)),
        "message": task.get("message") or task.get("result_message") or ("video tagging completed" if is_success else "video tagging failed"),
        "error_detail": task.get("error_detail") or task.get("error_message") or "",
    }


def validate_blogger_tagging_request(body):
    if not str(body.get("tiktok_blogger_id") or "").strip():
        return False, {"code": 4001, "message": "tiktok_blogger_id is required"}, None
    if not str(body.get("callback_url") or "").strip():
        return False, {"code": 4002, "message": "callback_url is required"}, None
    try:
        uuid.UUID(str(body["tiktok_blogger_id"]))
    except Exception:
        return False, {"code": 4003, "message": "tiktok_blogger_id must be a UUID"}, None
    callback_url = str(body["callback_url"]).strip()
    if not callback_url.startswith(("http://", "https://")):
        return False, {"code": 4004, "message": "callback_url must start with http:// or https://"}, None
    try:
        min_video_count = int(body.get("min_video_count") or current_blogger_min_video_count())
    except Exception:
        min_video_count = current_blogger_min_video_count()
    min_video_count = max(1, min(min_video_count, 50))
    return True, None, {
        "tiktok_blogger_id": str(body["tiktok_blogger_id"]).strip(),
        "callback_url": callback_url,
        "min_video_count": min_video_count,
    }


def blogger_tagging_table_sql():
    return """
create table if not exists public.blogger_tagging_results (
  id uuid primary key,
  tiktok_blogger_id uuid not null unique,
  status varchar not null,
  result_code integer,
  result_message text,
  error_code varchar,
  error_message text,
  min_video_count integer not null default 15,
  available_video_count integer not null default 0,
  usable_video_count integer not null default 0,
  successful_video_count integer not null default 0,
  failed_video_count integer not null default 0,
  submitted_video_count integer not null default 0,
  selected_video_ids uuid[],
  video_task_ids uuid[],
  account_personal_tags jsonb,
  account_style_vector jsonb,
  account_style_signature jsonb,
  account_one_sentence_summary text,
  aggregated_social_identity jsonb,
  aggregated_occasion jsonb,
  raw_outputs jsonb,
  callback_url text not null,
  callback_status varchar,
  callback_response_code integer,
  callback_response_body text,
  callback_attempts integer not null default 0,
  last_callback_at timestamptz,
  worker_id varchar,
  lock_until timestamptz,
  attempts integer not null default 0,
  next_retry_at timestamptz,
  created_at timestamptz not null default now(),
  updated_at timestamptz not null default now(),
  started_at timestamptz,
  finished_at timestamptz
);
"""


def blogger_tagging_migration_sql():
    return """
alter table public.blogger_tagging_results add column if not exists worker_id varchar;
alter table public.blogger_tagging_results add column if not exists lock_until timestamptz;
alter table public.blogger_tagging_results add column if not exists attempts integer not null default 0;
alter table public.blogger_tagging_results add column if not exists next_retry_at timestamptz;
alter table public.blogger_tagging_results add column if not exists account_one_sentence_summary text;
"""


def video_tagging_migration_sql():
    return """
alter table public.video_tagging_results add column if not exists gcs_url text;
do $$
begin
  if exists (
    select 1 from information_schema.columns
    where table_schema = 'public'
      and table_name = 'video_tagging_results'
      and column_name = 'local_video_url'
  ) then
    alter table public.video_tagging_results alter column local_video_url drop not null;
    update public.video_tagging_results
    set gcs_url = local_video_url
    where gcs_url is null and local_video_url is not null;
  end if;
end $$;
alter table public.video_tagging_results add column if not exists worker_id varchar;
alter table public.video_tagging_results add column if not exists lock_until timestamptz;
alter table public.video_tagging_results add column if not exists attempts integer not null default 0;
alter table public.video_tagging_results add column if not exists next_retry_at timestamptz;
alter table public.video_tagging_results add column if not exists source_type varchar not null default 'direct';
alter table public.video_tagging_results add column if not exists source_blogger_task_id uuid;
alter table public.video_tagging_results add column if not exists source_tiktok_blogger_id uuid;
"""


def build_blogger_tagging_callback_payload(task):
    status = task.get("status")
    is_success = status == "success"
    return {
        "event": "blogger_tagging.completed" if is_success else "blogger_tagging.failed",
        "tiktok_blogger_id": str(task.get("tiktok_blogger_id")),
        "task_id": str(task.get("id")),
        "status": status,
        "code": int(task.get("code") if task.get("code") is not None else task.get("result_code") or (0 if is_success else 5000)),
        "message": task.get("message") or task.get("result_message") or ("blogger tagging completed" if is_success else "blogger tagging failed"),
        "error_detail": task.get("error_detail") or task.get("error_message") or "",
    }


def merge_blogger_personal_tags(llm_tags, classification_summary):
    merged = {
        "basic_demographics": llm_tags.get("basic_demographics") or {},
        "consumption_tier": llm_tags.get("consumption_tier") or "",
        "temperament_psychology": llm_tags.get("temperament_psychology") or "",
        "social_identity": (classification_summary or {}).get("social_identity", {}).get("final_label") or "",
        "occasion": (classification_summary or {}).get("occasion", {}).get("final_label") or "",
    }
    if llm_tags.get("confidence"):
        merged["confidence"] = llm_tags["confidence"]
    return merged


def get_asyncpg_module():
    try:
        import asyncpg

        return asyncpg
    except ImportError:
        fallback = Path("/tmp/codex_pgdeps")
        if fallback.exists() and str(fallback) not in sys.path:
            sys.path.insert(0, str(fallback))
        import asyncpg

        return asyncpg


def db_dsn():
    if not DATABASE_URL:
        raise RuntimeError("Missing DATABASE_URL")
    return DATABASE_URL.replace("postgresql+asyncpg://", "postgresql://", 1)


async def db_connect():
    asyncpg = get_asyncpg_module()
    timeout = float(os.getenv("DB_CONNECT_TIMEOUT", "8"))
    try:
        return await asyncpg.connect(db_dsn(), timeout=timeout)
    except (TimeoutError, OSError) as exc:
        raise RuntimeError("Database connection failed or timed out") from exc


def db_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def row_to_task(row):
    if not row:
        return None
    data = dict(row)
    if not data.get("gcs_url") and data.get("local_video_url"):
        data["gcs_url"] = data.get("local_video_url")
    data.pop("local_video_url", None)
    for key in (
        "video_description_unit",
        "personal_tags",
        "style_vector",
        "style_signature",
        "account_personal_tags",
        "account_style_vector",
        "account_style_signature",
        "aggregated_social_identity",
        "aggregated_occasion",
        "raw_outputs",
    ):
        if isinstance(data.get(key), str):
            try:
                data[key] = json.loads(data[key])
            except Exception:
                pass
    for key in (
        "id",
        "video_id",
        "tiktok_blogger_id",
        "source_blogger_task_id",
        "source_tiktok_blogger_id",
        "selected_video_ids",
        "video_task_ids",
        "created_at",
        "updated_at",
        "started_at",
        "finished_at",
        "last_callback_at",
        "lock_until",
        "next_retry_at",
    ):
        if data.get(key) is not None:
            if isinstance(data[key], (list, tuple)):
                data[key] = [str(item) for item in data[key]]
            else:
                data[key] = str(data[key])
    return data


async def ensure_video_tagging_table_async():
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        await conn.execute(blogger_tagging_table_sql())
        await conn.execute(video_tagging_migration_sql())
        await conn.execute(blogger_tagging_migration_sql())
    finally:
        await conn.close()


def ensure_video_tagging_table():
    return asyncio.run(ensure_video_tagging_table_async())


async def ensure_blogger_tagging_table_async():
    conn = await db_connect()
    try:
        await conn.execute(blogger_tagging_table_sql())
        await conn.execute(blogger_tagging_migration_sql())
    finally:
        await conn.close()


async def ensure_all_tagging_tables_async():
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        await conn.execute(blogger_tagging_table_sql())
        await conn.execute(video_tagging_migration_sql())
        await conn.execute(blogger_tagging_migration_sql())
    finally:
        await conn.close()


def ensure_all_tagging_tables():
    return asyncio.run(ensure_all_tagging_tables_async())


async def submit_video_tagging_task_async(body):
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        existing = await conn.fetchrow(
            """
            select * from public.video_tagging_results
            where video_id = $1::uuid
            """,
            body["video_id"],
        )
        if existing and existing["status"] == "success":
            task = row_to_task(existing)
            return {"code": 0, "message": "already_processed", "task": task, "start_worker": False}
        if existing and existing["status"] in ("pending", "running"):
            task = row_to_task(existing)
            return {"code": 0, "message": "already_running", "task": task, "start_worker": False}
        task_id = str(existing["id"]) if existing else str(uuid.uuid4())
        if existing:
            existing_data = dict(existing)
            source_type = str(body.get("source_type") or existing_data.get("source_type") or "direct")
            source_blogger_task_id = body.get("source_blogger_task_id") or existing_data.get("source_blogger_task_id")
            source_tiktok_blogger_id = body.get("source_tiktok_blogger_id") or existing_data.get("source_tiktok_blogger_id")
            row = await conn.fetchrow(
                """
                update public.video_tagging_results
                set gcs_url = $2,
                    description = $3,
                    callback_url = $4,
                    status = 'pending',
                    result_code = null,
                    result_message = 'received',
                    error_code = null,
                    error_message = null,
                    callback_status = null,
                    callback_response_code = null,
                    callback_response_body = null,
                    callback_attempts = 0,
                    last_callback_at = null,
                    source_type = $5,
                    source_blogger_task_id = $6::uuid,
                    source_tiktok_blogger_id = $7::uuid,
                    updated_at = now(),
                    started_at = null,
                    finished_at = null
                where id = $1::uuid
                returning *
                """,
                task_id,
                body["gcs_url"],
                body["description"],
                body["callback_url"],
                source_type,
                source_blogger_task_id,
                source_tiktok_blogger_id,
            )
        else:
            row = await conn.fetchrow(
                """
                insert into public.video_tagging_results (
                    id, video_id, gcs_url, description, callback_url,
                    status, result_code, result_message,
                    source_type, source_blogger_task_id, source_tiktok_blogger_id
                )
                values ($1::uuid, $2::uuid, $3, $4, $5, 'pending', 0, 'received', $6, $7::uuid, $8::uuid)
                returning *
                """,
                task_id,
                body["video_id"],
                body["gcs_url"],
                body["description"],
                body["callback_url"],
                str(body.get("source_type") or "direct"),
                body.get("source_blogger_task_id"),
                body.get("source_tiktok_blogger_id"),
            )
        return {"code": 0, "message": "received", "task": row_to_task(row), "start_worker": True}
    finally:
        await conn.close()


def submit_video_tagging_task(body):
    return asyncio.run(submit_video_tagging_task_async(body))


async def fetch_video_tagging_task_async(video_id):
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        return row_to_task(
            await conn.fetchrow(
                "select * from public.video_tagging_results where video_id = $1::uuid",
                video_id,
            )
        )
    finally:
        await conn.close()


def fetch_video_tagging_task(video_id):
    return asyncio.run(fetch_video_tagging_task_async(video_id))


async def fetch_video_signed_url_async(video_id):
    task = await fetch_video_tagging_task_async(video_id)
    if not task:
        return None
    signed_url = ensure_fresh_gcs_signed_url(task.get("gcs_url") or "")
    return {
        "video_id": task.get("video_id"),
        "gcs_url": task.get("gcs_url"),
        "signed_gcs_url": signed_url,
        "fresh": is_gcs_signed_url_fresh(signed_url),
    }


def fetch_video_signed_url(video_id):
    return asyncio.run(fetch_video_signed_url_async(video_id))


async def fetch_video_tagging_task_by_id_async(task_id):
    conn = await db_connect()
    try:
        return row_to_task(
            await conn.fetchrow(
                "select * from public.video_tagging_results where id = $1::uuid",
                task_id,
            )
        )
    finally:
        await conn.close()


async def list_video_tagging_tasks_async(status=None, limit=100, date=None, include_blogger=False):
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        await conn.execute(blogger_tagging_table_sql())
        limit = max(1, min(int(limit), 500))
        start_at, end_at = beijing_day_to_utc_range(date)
        where = []
        values = []
        if not include_blogger:
            where.append("(coalesce(source_type, 'direct') <> 'blogger')")
            where.append(
                """not exists (
                    select 1
                    from public.blogger_tagging_results b
                    where public.video_tagging_results.id = any(b.video_task_ids)
                )"""
            )
        if status:
            values.append(status)
            where.append(f"status = ${len(values)}")
        if start_at and end_at:
            values.extend([start_at, end_at])
            where.append(f"created_at >= ${len(values) - 1} and created_at < ${len(values)}")
        values.append(limit)
        sql = f"""
            select * from public.video_tagging_results
            {"where " + " and ".join(where) if where else ""}
            order by created_at desc
            limit ${len(values)}
        """
        rows = await conn.fetch(sql, *values)
        return [row_to_task(row) for row in rows]
    finally:
        await conn.close()


def list_video_tagging_tasks(status=None, limit=100, date=None, include_blogger=False):
    return asyncio.run(list_video_tagging_tasks_async(status, limit, date, include_blogger))


async def enrich_blogger_source_links(conn, tasks, video_url_limit=5):
    blogger_ids = []
    for task in tasks:
        blogger_id = task.get("tiktok_blogger_id")
        if blogger_id and blogger_id not in blogger_ids:
            blogger_ids.append(blogger_id)
    if not blogger_ids:
        return tasks

    blogger_rows = await conn.fetch(
        """
        select id, blogger_url
        from public.tiktok_bloggers
        where id = any($1::uuid[])
        """,
        blogger_ids,
    )
    blogger_urls = {str(row["id"]): row["blogger_url"] or "" for row in blogger_rows}

    video_rows = await conn.fetch(
        """
        select blogger_id, source_url, source_video_count
        from (
          select
            v.tiktok_blogger_id as blogger_id,
            coalesce(nullif(v.source_url, ''), nullif(v.video_url, ''), nullif(cv.video_url, '')) as source_url,
            count(*) over (partition by v.tiktok_blogger_id) as source_video_count,
            row_number() over (
              partition by v.tiktok_blogger_id
              order by v.publish_date desc nulls last, v.created_at desc
            ) as row_index
          from public.video_sources v
          left join lateral (
            select video_url
            from public.candidate_videos
            where video_source_id = v.id and nullif(video_url, '') is not null
            order by created_at desc
            limit 1
          ) cv on true
          where v.tiktok_blogger_id = any($1::uuid[])
            and coalesce(nullif(v.source_url, ''), nullif(v.video_url, ''), nullif(cv.video_url, '')) is not null
        ) linked
        where row_index <= $2
        order by blogger_id, row_index
        """,
        blogger_ids,
        max(1, int(video_url_limit)),
    )
    source_urls = {}
    source_counts = {}
    for row in video_rows:
        blogger_id = str(row["blogger_id"])
        source_urls.setdefault(blogger_id, []).append(row["source_url"])
        source_counts[blogger_id] = int(row["source_video_count"] or 0)

    for task in tasks:
        blogger_id = task.get("tiktok_blogger_id")
        task["blogger_url"] = blogger_urls.get(blogger_id, "")
        task["source_video_urls"] = source_urls.get(blogger_id, [])
        task["source_video_count"] = source_counts.get(blogger_id, len(task["source_video_urls"]))
    return tasks


async def list_blogger_video_tagging_tasks_async(tiktok_blogger_id, status=None, limit=100, date=None):
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        await conn.execute(blogger_tagging_table_sql())
        limit = max(1, min(int(limit), 500))
        start_at, end_at = beijing_day_to_utc_range(date)
        blogger_task = await conn.fetchrow(
            """
            select video_task_ids, selected_video_ids
            from public.blogger_tagging_results
            where tiktok_blogger_id = $1::uuid
            """,
            tiktok_blogger_id,
        )
        blogger = await conn.fetchrow(
            """
            select blogger_url
            from public.tiktok_bloggers
            where id = $1::uuid
            """,
            tiktok_blogger_id,
        )
        blogger_url = blogger["blogger_url"] if blogger and blogger["blogger_url"] else ""
        video_task_ids = list(blogger_task["video_task_ids"] or []) if blogger_task else []
        selected_video_ids = list(blogger_task["selected_video_ids"] or []) if blogger_task else []
        blogger_video_rows = await conn.fetch(
            "select id from public.video_sources where tiktok_blogger_id = $1::uuid",
            tiktok_blogger_id,
        )
        blogger_video_ids = [row["id"] for row in blogger_video_rows] + selected_video_ids
        values = [tiktok_blogger_id, video_task_ids, blogger_video_ids]
        where = [
            "(vt.source_tiktok_blogger_id = $1::uuid or vt.id = any($2::uuid[]) or vt.video_id = any($3::uuid[]))"
        ]
        if status:
            values.append(status)
            where.append(f"vt.status = ${len(values)}")
        if start_at and end_at:
            values.extend([start_at, end_at])
            where.append(f"vt.created_at >= ${len(values) - 1} and vt.created_at < ${len(values)}")
        values.append(limit)
        rows = await conn.fetch(
            f"""
            select
              vt.*,
              coalesce(nullif(v.source_url, ''), nullif(v.video_url, ''), nullif(cv.video_url, '')) as source_url
            from public.video_tagging_results vt
            left join public.video_sources v on v.id = vt.video_id
            left join lateral (
              select video_url
              from public.candidate_videos
              where video_source_id = v.id and nullif(video_url, '') is not null
              order by created_at desc
              limit 1
            ) cv on true
            where {" and ".join(where)}
            order by vt.created_at desc
            limit ${len(values)}
            """,
            *values,
        )
        tasks = [row_to_task(row) for row in rows]
        for task in tasks:
            task["blogger_url"] = blogger_url
        return tasks
    finally:
        await conn.close()


def list_blogger_video_tagging_tasks(tiktok_blogger_id, status=None, limit=100, date=None):
    return asyncio.run(list_blogger_video_tagging_tasks_async(tiktok_blogger_id, status, limit, date))


async def update_video_tagging_task_async(task_id, **fields):
    conn = await db_connect()
    try:
        allowed = {
            "status",
            "result_code",
            "result_message",
            "error_code",
            "error_message",
            "video_description_unit",
            "personal_tags",
            "style_vector",
            "style_signature",
            "raw_outputs",
            "callback_status",
            "callback_response_code",
            "callback_response_body",
            "callback_attempts",
            "worker_id",
            "lock_until",
            "attempts",
            "next_retry_at",
        }
        items = [(key, value) for key, value in fields.items() if key in allowed]
        if not items:
            return await fetch_video_tagging_task_by_id_async(task_id)
        assignments = []
        values = [task_id]
        for index, (key, value) in enumerate(items, start=2):
            cast = "::jsonb" if key in {
                "video_description_unit",
                "personal_tags",
                "style_vector",
                "style_signature",
                "raw_outputs",
            } else ""
            assignments.append(f"{key} = ${index}{cast}")
            values.append(db_value(value))
        assignments.append("updated_at = now()")
        if fields.get("status") == "running":
            assignments.append("started_at = now()")
        if fields.get("status") in ("success", "failed"):
            assignments.append("finished_at = now()")
            assignments.append("worker_id = null")
            assignments.append("lock_until = null")
        if "callback_status" in fields:
            assignments.append("last_callback_at = now()")
        sql = f"""
            update public.video_tagging_results
            set {", ".join(assignments)}
            where id = $1::uuid
            returning *
        """
        return row_to_task(await conn.fetchrow(sql, *values))
    finally:
        await conn.close()


def update_video_tagging_task(task_id, **fields):
    return asyncio.run(update_video_tagging_task_async(task_id, **fields))


async def fetch_blogger_profile_async(tiktok_blogger_id):
    conn = await db_connect()
    try:
        row = await conn.fetchrow(
            "select signature from public.tiktok_bloggers where id = $1::uuid",
            tiktok_blogger_id,
        )
        if not row or not row["signature"]:
            return ""
        return str(row["signature"]).strip()
    finally:
        await conn.close()


def fetch_blogger_profile(tiktok_blogger_id):
    return asyncio.run(fetch_blogger_profile_async(tiktok_blogger_id))


async def fetch_blogger_videos_async(conn, tiktok_blogger_id):
    rows = await conn.fetch(
        """
        select
          v.id as video_id,
          v.local_gcs_video_url as gcs_url,
          coalesce(v.video_desc, v.video_title, cv.video_title, '') as description,
          v.created_at,
          v.publish_date
        from public.video_sources v
        left join public.candidate_videos cv on cv.video_source_id = v.id
        where v.tiktok_blogger_id = $1::uuid
        order by v.publish_date desc nulls last, v.created_at desc
        """,
        tiktok_blogger_id,
    )
    videos = []
    seen = set()
    for row in rows:
        video_id = str(row["video_id"])
        if video_id in seen:
            continue
        seen.add(video_id)
        if not row["gcs_url"] or not row["description"]:
            continue
        videos.append(
            {
                "video_id": video_id,
                "gcs_url": row["gcs_url"],
                "description": row["description"],
            }
        )
    return videos


async def fetch_blogger_videos_db_async(tiktok_blogger_id):
    conn = await db_connect()
    try:
        return await fetch_blogger_videos_async(conn, tiktok_blogger_id)
    finally:
        await conn.close()


def fetch_blogger_videos_db(tiktok_blogger_id):
    return asyncio.run(fetch_blogger_videos_db_async(tiktok_blogger_id))


async def fetch_video_tagging_tasks_by_ids_async(video_ids):
    if not video_ids:
        return []
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        rows = await conn.fetch(
            """
            select *
            from public.video_tagging_results
            where video_id = any($1::uuid[])
            order by created_at desc
            """,
            video_ids,
        )
        return [row_to_task(row) for row in rows]
    finally:
        await conn.close()


def fetch_video_tagging_tasks_by_ids(video_ids):
    return asyncio.run(fetch_video_tagging_tasks_by_ids_async(video_ids))


def video_task_claim_sql():
    return """
with next_task as (
  select id
  from public.video_tagging_results
  where (
      status = 'pending'
      or (status = 'running' and (lock_until is null or lock_until < now()))
    )
    and (next_retry_at is null or next_retry_at <= now())
  order by created_at asc
  for update skip locked
  limit 1
)
update public.video_tagging_results
set status = 'running',
    result_message = 'running',
    worker_id = $1,
    lock_until = now() + ($2 * interval '1 second'),
    attempts = attempts + 1,
    updated_at = now(),
    started_at = coalesce(started_at, now())
where id = (select id from next_task)
returning *;
"""


def blogger_task_claim_sql():
    return """
with next_task as (
  select id
  from public.blogger_tagging_results
  where (
      status = 'pending'
      or (
        status = 'waiting_videos'
        and (next_retry_at is null or next_retry_at <= now())
      )
      or (
        status in ('checking_videos', 'aggregating')
        and (lock_until is null or lock_until < now())
      )
    )
  order by created_at asc
  for update skip locked
  limit 1
)
update public.blogger_tagging_results
set status = 'checking_videos',
    result_message = 'checking videos',
    worker_id = $1,
    lock_until = now() + ($2 * interval '1 second'),
    attempts = attempts + 1,
    updated_at = now(),
    started_at = coalesce(started_at, now())
where id = (select id from next_task)
returning *;
"""


async def claim_next_video_tagging_task_async(worker_id):
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        row = await conn.fetchrow(video_task_claim_sql(), worker_id, DEFAULT_TASK_LOCK_SECONDS)
        return row_to_task(row)
    finally:
        await conn.close()


def claim_next_video_tagging_task(worker_id):
    return asyncio.run(claim_next_video_tagging_task_async(worker_id))


async def claim_next_blogger_tagging_task_async(worker_id):
    conn = await db_connect()
    try:
        await conn.execute(blogger_tagging_table_sql())
        row = await conn.fetchrow(blogger_task_claim_sql(), worker_id, DEFAULT_TASK_LOCK_SECONDS)
        return row_to_task(row)
    finally:
        await conn.close()


def claim_next_blogger_tagging_task(worker_id):
    return asyncio.run(claim_next_blogger_tagging_task_async(worker_id))


async def submit_blogger_tagging_task_async(body):
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        await conn.execute(blogger_tagging_table_sql())
        exists = await conn.fetchval(
            "select exists(select 1 from public.tiktok_bloggers where id = $1::uuid)",
            body["tiktok_blogger_id"],
        )
        if not exists:
            return {"code": 4201, "message": "blogger not found", "task": None, "start_worker": False}

        videos = await fetch_blogger_videos_async(conn, body["tiktok_blogger_id"])
        available = len(videos)
        if available < body["min_video_count"]:
            return {
                "code": 4202,
                "message": "insufficient videos",
                "available_video_count": available,
                "required_video_count": body["min_video_count"],
                "task": None,
                "start_worker": False,
            }

        existing = await conn.fetchrow(
            """
            select * from public.blogger_tagging_results
            where tiktok_blogger_id = $1::uuid
            """,
            body["tiktok_blogger_id"],
        )
        if existing and existing["status"] == "success":
            return {"code": 0, "message": "already_processed", "task": row_to_task(existing), "start_worker": False}
        if existing and existing["status"] in ("pending", "checking_videos", "waiting_videos", "aggregating"):
            return {"code": 0, "message": "already_running", "task": row_to_task(existing), "start_worker": False}

        task_id = str(existing["id"]) if existing else str(uuid.uuid4())
        if existing:
            row = await conn.fetchrow(
                """
                update public.blogger_tagging_results
                set callback_url = $2,
                    min_video_count = $3,
                    available_video_count = $4,
                    usable_video_count = $4,
                    status = 'pending',
                    result_code = 0,
                    result_message = 'received',
                    error_code = null,
                    error_message = null,
                    callback_status = null,
                    callback_response_code = null,
                    callback_response_body = null,
                    callback_attempts = 0,
                    worker_id = null,
                    lock_until = null,
                    next_retry_at = null,
                    updated_at = now(),
                    started_at = null,
                    finished_at = null
                where id = $1::uuid
                returning *
                """,
                task_id,
                body["callback_url"],
                body["min_video_count"],
                available,
            )
        else:
            row = await conn.fetchrow(
                """
                insert into public.blogger_tagging_results (
                    id, tiktok_blogger_id, callback_url, min_video_count,
                    available_video_count, usable_video_count, status,
                    result_code, result_message
                )
                values ($1::uuid, $2::uuid, $3, $4, $5, $5, 'pending', 0, 'received')
                returning *
                """,
                task_id,
                body["tiktok_blogger_id"],
                body["callback_url"],
                body["min_video_count"],
                available,
            )
        return {"code": 0, "message": "received", "task": row_to_task(row), "start_worker": True}
    finally:
        await conn.close()


def submit_blogger_tagging_task(body):
    return asyncio.run(submit_blogger_tagging_task_async(body))


async def fetch_blogger_tagging_task_async(tiktok_blogger_id):
    conn = await db_connect()
    try:
        await conn.execute(blogger_tagging_table_sql())
        return row_to_task(
            await conn.fetchrow(
                "select * from public.blogger_tagging_results where tiktok_blogger_id = $1::uuid",
                tiktok_blogger_id,
            )
        )
    finally:
        await conn.close()


def fetch_blogger_tagging_task(tiktok_blogger_id):
    return asyncio.run(fetch_blogger_tagging_task_async(tiktok_blogger_id))


async def fetch_blogger_tagging_task_by_id_async(task_id):
    conn = await db_connect()
    try:
        return row_to_task(
            await conn.fetchrow(
                "select * from public.blogger_tagging_results where id = $1::uuid",
                task_id,
            )
        )
    finally:
        await conn.close()


async def list_blogger_tagging_tasks_async(status=None, limit=100, date=None):
    conn = await db_connect()
    try:
        await conn.execute(blogger_tagging_table_sql())
        limit = max(1, min(int(limit), 500))
        start_at, end_at = beijing_day_to_utc_range(date)
        where = []
        values = []
        if status:
            values.append(status)
            where.append(f"status = ${len(values)}")
        if start_at and end_at:
            values.extend([start_at, end_at])
            where.append(f"created_at >= ${len(values) - 1} and created_at < ${len(values)}")
        values.append(limit)
        sql = f"""
            select * from public.blogger_tagging_results
            {"where " + " and ".join(where) if where else ""}
            order by created_at desc
            limit ${len(values)}
        """
        rows = await conn.fetch(sql, *values)
        tasks = [row_to_task(row) for row in rows]
        return await enrich_blogger_source_links(conn, tasks)
    finally:
        await conn.close()


def list_blogger_tagging_tasks(status=None, limit=100, date=None):
    return asyncio.run(list_blogger_tagging_tasks_async(status, limit, date))


async def update_blogger_tagging_task_async(task_id, **fields):
    conn = await db_connect()
    try:
        allowed = {
            "status",
            "result_code",
            "result_message",
            "error_code",
            "error_message",
            "available_video_count",
            "usable_video_count",
            "successful_video_count",
            "failed_video_count",
            "submitted_video_count",
            "selected_video_ids",
            "video_task_ids",
            "account_personal_tags",
            "account_style_vector",
            "account_style_signature",
            "account_one_sentence_summary",
            "aggregated_social_identity",
            "aggregated_occasion",
            "raw_outputs",
            "callback_status",
            "callback_response_code",
            "callback_response_body",
            "callback_attempts",
            "worker_id",
            "lock_until",
            "attempts",
            "next_retry_at",
        }
        items = [(key, value) for key, value in fields.items() if key in allowed]
        if not items:
            return await fetch_blogger_tagging_task_by_id_async(task_id)
        assignments = []
        values = [task_id]
        json_keys = {
            "account_personal_tags",
            "account_style_vector",
            "account_style_signature",
            "aggregated_social_identity",
            "aggregated_occasion",
            "raw_outputs",
        }
        uuid_array_keys = {"selected_video_ids", "video_task_ids"}
        for index, (key, value) in enumerate(items, start=2):
            if key in json_keys:
                cast = "::jsonb"
            elif key in uuid_array_keys:
                cast = "::uuid[]"
            else:
                cast = ""
            assignments.append(f"{key} = ${index}{cast}")
            values.append(value if key in uuid_array_keys else db_value(value))
        assignments.append("updated_at = now()")
        if fields.get("status") in ("checking_videos", "waiting_videos", "aggregating"):
            assignments.append("started_at = coalesce(started_at, now())")
        if fields.get("status") in ("success", "failed"):
            assignments.append("finished_at = now()")
            assignments.append("worker_id = null")
            assignments.append("lock_until = null")
        if "callback_status" in fields:
            assignments.append("last_callback_at = now()")
        sql = f"""
            update public.blogger_tagging_results
            set {", ".join(assignments)}
            where id = $1::uuid
            returning *
        """
        return row_to_task(await conn.fetchrow(sql, *values))
    finally:
        await conn.close()


def update_blogger_tagging_task(task_id, **fields):
    return asyncio.run(update_blogger_tagging_task_async(task_id, **fields))


def load_default_prompt(number):
    config_prompts = read_service_config().get("prompts")
    if isinstance(config_prompts, dict) and config_prompts.get(str(number)):
        return config_prompts[str(number)]
    path = STATIC_DIR / "prompts.js"
    if not path.exists():
        return ""
    text = path.read_text(errors="ignore")
    match = re.search(
        rf"window\.DEFAULT_PROMPT_{number}\s*=\s*`(.*?)`;",
        text,
        re.S,
    )
    return match.group(1) if match else ""


def require_parsed(stage_name, result):
    if result.get("error"):
        raise RuntimeError(result["error"])
    if not isinstance(result.get("parsed"), dict):
        raise ValueError(f"{stage_name} returned non-JSON result")
    return result["parsed"]


def video_tagging_error_code(exc):
    text = str(exc)
    if "Missing EVOLINK_API_KEY" in text:
        return 5001
    if "returned non-JSON" in text:
        return 5004
    if "timed out" in text.lower() or "timeout" in text.lower():
        return 5007
    if "Video API HTTP" in text:
        return 5002
    return 5001


def blogger_account_prompt():
    return """你是一个 TikTok 博主账号级 Personal Tags 打标模型。

我会提供同一个 TikTok 博主的多个 video_description_unit。
你只能根据这些单视频描述单元，判断这个账号稳定呈现的 3 类标签：
1. 基础人口
2. 消费层级
3. 气质心理

不要输出 social_identity。
不要输出 occasion。
不要输出 style。
不要输出账号定位总结。
不要输出证据摘要。
不要输出 schema 之外的字段。

如果信息不足，必须选择“无明显 / 无明确”类标签。
只有多条视频反复稳定出现的信号，才作为账号级最终标签。

基础人口标签：
性别 / 性向呈现：男 / 女 / Gay / Les / Trans / 无明显
年龄段：18-24 / 25-34 / 35-44 / 45+ / 无明显
族裔视觉：非裔 / 亚裔 / 拉美裔 / 白人 / 印度人 / 无明显
身材类型：Super Fat / Large / Plump / Normal / Skinny / Muscular / 无明显
身高感：Short / Normal / Tall / 无明显
特殊 Body 部位：胸 / 手 / 脚 / 腿 / 屁股 / 头发 / 背 / belly / 无明显特殊 Body 部位

消费层级：
高性价比：$0 - $60 / 中端通勤：$60 - $180 / 轻奢：$180 - $500 / 奢侈品：$500+ / 无明显消费层级 / Remix 消费型

气质心理：
温柔亲和型 / 自信性感型 / 冷感高级型 / 活力阳光型 / 搞笑混乱型 / 安静内向型 / 精英利落型 / 叛逆个性型 / 无明显气质类型 / Remix 气质型

请只输出 JSON：
{
  "basic_demographics": {
    "gender_or_sexuality_presentation": "",
    "age_range": "",
    "visual_ethnicity": "",
    "body_type": "",
    "height_impression": "",
    "special_body_parts": []
  },
  "consumption_tier": "",
  "temperament_psychology": "",
  "confidence": {
    "basic_demographics": "",
    "consumption_tier": "",
    "temperament_psychology": ""
  }
}"""


def build_blogger_account_prompt(units):
    return (
        f"{blogger_account_prompt()}\n\n"
        f"以下是该账号的 {len(units)} 个 video_description_unit：\n"
        f"{json.dumps(units, ensure_ascii=False, indent=2)}"
    )


def analyze_blogger_lite_account(units):
    try:
        result = call_evolink_text([{"role": "user", "content": build_blogger_account_prompt(units)}])
        text = extract_text(result)
        return {"raw_text": text, "parsed": parse_json_text(text), "error": ""}
    except Exception as exc:
        return {"raw_text": "", "parsed": None, "error": str(exc)}


def build_blogger_one_sentence_summary_prompt(base_prompt, units, blogger_profile=""):
    profile = str(blogger_profile or "").strip() or "无"
    return (
        f"{base_prompt}\n\n"
        f"以下是该 TikTok 博主主页 profile/bio 文案。它可能包含地点、职业、身份、内容定位、联系方式或自我介绍；"
        f"如果它提供了重要且可信的信息，请优先用于判断这个人是谁和账号定位，但不要编造 profile 里没有的信息：\n"
        f"{profile}\n\n"
        f"以下是该账号的 {len(units)} 个 video_description_unit：\n"
        f"{json.dumps(units, ensure_ascii=False, indent=2)}"
    )


def analyze_blogger_one_sentence_summary(prompt, units, blogger_profile=""):
    try:
        result = call_evolink_text(
            [{"role": "user", "content": build_blogger_one_sentence_summary_prompt(prompt, units, blogger_profile)}]
        )
        text = extract_text(result)
        return {"raw_text": text, "parsed": parse_json_text(text), "error": ""}
    except Exception as exc:
        return {"raw_text": "", "parsed": None, "error": str(exc)}


def extract_account_one_sentence_summary(result):
    parsed = result.get("parsed") if isinstance(result, dict) else None
    if not isinstance(parsed, dict):
        return ""
    value = parsed.get("account_one_sentence_summary")
    return value.strip() if isinstance(value, str) else ""


def video_result_to_classification(video_result):
    return {"parsed": video_result.get("personal_tags") or {}, "error": ""}


def video_result_to_style_vector(video_result):
    return {"parsed": video_result.get("style_vector") or {}, "error": ""}


def video_result_to_style_signature(video_result):
    return {"parsed": video_result.get("style_signature") or {}, "error": ""}


def submit_internal_video_task(video, callback_url, blogger_task_id=None, tiktok_blogger_id=None):
    return submit_video_tagging_task(
        {
            "video_id": video["video_id"],
            "gcs_url": video["gcs_url"],
            "description": video["description"],
            "callback_url": callback_url,
            "source_type": "blogger",
            "source_blogger_task_id": blogger_task_id,
            "source_tiktok_blogger_id": tiktok_blogger_id,
        }
    )


def internal_callback_url():
    return os.getenv(
        "INTERNAL_CALLBACK_URL",
        f"http://127.0.0.1:{os.getenv('PORT', '4190')}/api/internal/tagging-callback",
    )


def run_blogger_tagging_task(task_id):
    task = None
    try:
        task = asyncio.run(fetch_blogger_tagging_task_by_id_async(task_id))
        if not task:
            return
        task = update_blogger_tagging_task(
            task_id,
            status="checking_videos",
            result_message="checking videos",
        )
        videos = fetch_blogger_videos_db(task["tiktok_blogger_id"])

        min_count = int(task.get("min_video_count") or 15)
        if len(videos) < min_count:
            task = update_blogger_tagging_task(
                task_id,
                status="failed",
                result_code=4202,
                result_message="insufficient videos",
                error_code="4202",
                error_message=f"available videos {len(videos)} < required {min_count}",
                available_video_count=len(videos),
                usable_video_count=len(videos),
            )
            send_blogger_tagging_callback(task)
            return

        video_ids = [video["video_id"] for video in videos]
        rows = fetch_video_tagging_tasks_by_ids(video_ids)
        by_video = {str(row["video_id"]): row for row in rows}
        success = [by_video[video_id] for video_id in video_ids if by_video.get(video_id, {}).get("status") == "success"]
        failed_count = sum(1 for video_id in video_ids if by_video.get(video_id, {}).get("status") == "failed")

        submitted = 0
        video_task_ids = []
        if len(success) < min_count:
            needed = min_count - len(success)
            missing = [
                video
                for video in videos
                if by_video.get(video["video_id"], {}).get("status") not in ("success", "pending", "running")
            ][:needed]
            for video in missing:
                result = submit_internal_video_task(
                    video,
                    internal_callback_url(),
                    task["id"],
                    task["tiktok_blogger_id"],
                )
                if result.get("task"):
                    video_task_ids.append(result["task"]["id"])
                    submitted += 1
            task = update_blogger_tagging_task(
                task_id,
                status="waiting_videos",
                result_code=0,
                result_message="waiting for video tagging",
                available_video_count=len(videos),
                usable_video_count=len(videos),
                successful_video_count=len(success),
                failed_video_count=failed_count,
                submitted_video_count=submitted,
                video_task_ids=video_task_ids,
            )
            schedule_blogger_recheck(task_id)
            return

        selected = success[:min_count]
        task = update_blogger_tagging_task(
            task_id,
            status="aggregating",
            result_message="aggregating blogger result",
            available_video_count=len(videos),
            usable_video_count=len(videos),
            successful_video_count=len(success),
            failed_video_count=failed_count,
            selected_video_ids=[item["video_id"] for item in selected],
        )

        units = [item.get("video_description_unit") for item in selected if item.get("video_description_unit")]
        blogger_profile = fetch_blogger_profile(task["tiktok_blogger_id"])
        account_result = analyze_blogger_lite_account(units)
        account_parsed = require_parsed("blogger_account", account_result)
        summary_result = analyze_blogger_one_sentence_summary(load_default_prompt(6), units, blogger_profile)
        account_one_sentence_summary = extract_account_one_sentence_summary(summary_result)
        classification_summary = aggregate_classifications(
            [video_result_to_classification(item) for item in selected]
        )
        style_summary = aggregate_style_results(
            [video_result_to_style_vector(item) for item in selected],
            [video_result_to_style_signature(item) for item in selected],
        )
        account_personal_tags = merge_blogger_personal_tags(account_parsed, classification_summary)
        task = update_blogger_tagging_task(
            task_id,
            status="success",
            result_code=0,
            result_message="blogger tagging completed",
            error_code=None,
            error_message=None,
            account_personal_tags=account_personal_tags,
            account_style_vector=style_summary.get("average_style_vector") or {},
            account_style_signature=style_summary.get("account_style_signature") or {},
            account_one_sentence_summary=account_one_sentence_summary,
            aggregated_social_identity=classification_summary.get("social_identity") or {},
            aggregated_occasion=classification_summary.get("occasion") or {},
            raw_outputs={
                "account_result": account_result,
                "one_sentence_summary": summary_result,
                "blogger_profile": blogger_profile,
                "style_summary": style_summary,
            },
        )
        send_blogger_tagging_callback(task)
    except Exception as exc:
        code = 5104 if "blogger_account" in str(exc) else 5000
        task = update_blogger_tagging_task(
            task_id,
            status="failed",
            result_code=code,
            result_message=str(exc),
            error_code=str(code),
            error_message=str(exc),
        )
        send_blogger_tagging_callback(task)


def send_blogger_tagging_callback(task):
    payload = build_blogger_tagging_callback_payload(task)
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    attempts = int(task.get("callback_attempts") or 0) + 1
    try:
        request = urllib.request.Request(
            task["callback_url"],
            data=data,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=15) as response:
            body = response.read().decode("utf-8", errors="ignore")
            update_blogger_tagging_task(
                task["id"],
                callback_status="success" if 200 <= response.status < 300 else "failed",
                callback_response_code=response.status,
                callback_response_body=body[:2000],
                callback_attempts=attempts,
            )
    except Exception as exc:
        update_blogger_tagging_task(
            task["id"],
            callback_status="failed",
            callback_response_code=0,
            callback_response_body=str(exc)[:2000],
            callback_attempts=attempts,
        )


def schedule_blogger_recheck(task_id, delay=30):
    update_blogger_tagging_task(
        task_id,
        next_retry_at=datetime.now(timezone.utc) + timedelta(seconds=delay),
    )


def recover_incomplete_tagging_tasks():
    try:
        ensure_all_tagging_tables()
        asyncio.run(recover_incomplete_tagging_tasks_async())
    except Exception as exc:
        print(f"Recover tagging tasks failed: {exc}")


async def recover_incomplete_tagging_tasks_async():
    conn = await db_connect()
    try:
        await conn.execute(video_tagging_table_sql())
        await conn.execute(blogger_tagging_table_sql())
        await conn.execute(
            """
            update public.video_tagging_results
            set status = 'pending',
                result_message = 'recovered after restart',
                worker_id = null,
                lock_until = null,
                updated_at = now()
            where status = 'running'
            """
        )
        await conn.execute(
            """
            update public.blogger_tagging_results
            set status = 'pending',
                result_message = 'recovered after restart',
                worker_id = null,
                lock_until = null,
                updated_at = now()
            where status in ('running', 'checking_videos', 'aggregating')
            """
        )
        await conn.execute(
            """
            update public.blogger_tagging_results
            set next_retry_at = coalesce(next_retry_at, now()),
                worker_id = null,
                lock_until = null,
                updated_at = now()
            where status = 'waiting_videos'
            """
        )
    finally:
        await conn.close()


def run_video_tagging_task(task_id):
    task = None
    raw_outputs = {}
    try:
        task = asyncio.run(fetch_video_tagging_task_by_id_async(task_id))
        if not task:
            return
        update_video_tagging_task(task_id, status="running", result_message="running")
        gcs_url = ensure_fresh_gcs_signed_url(task["gcs_url"])
        video = {
            "source_url": task["gcs_url"],
            "file_uri": gcs_url,
            "caption": task["description"],
            "hashtag": "",
        }
        api_gate = threading.BoundedSemaphore(1)
        prompt1 = load_default_prompt(1)
        prompt3 = load_default_prompt(3)
        prompt4 = load_default_prompt(4)
        prompt5 = load_default_prompt(5)

        unit = analyze_video_unit(prompt1, video, 1, api_gate)
        raw_outputs["video_description_unit"] = unit
        unit_parsed = require_parsed("video_description_unit", unit)

        classification = analyze_video_classification(prompt3, video, unit, api_gate)
        raw_outputs["personal_tags"] = classification
        personal_tags = require_parsed("personal_tags", classification)

        style_vector = analyze_video_style_vector(prompt4, video, unit, api_gate)
        raw_outputs["style_vector"] = style_vector
        style_vector_parsed = require_parsed("style_vector", style_vector)

        style_signature = analyze_video_style_signature(prompt5, video, unit, style_vector, api_gate)
        raw_outputs["style_signature"] = style_signature
        style_signature_parsed = require_parsed("style_signature", style_signature)

        task = update_video_tagging_task(
            task_id,
            status="success",
            result_code=0,
            result_message="video tagging completed",
            error_code=None,
            error_message=None,
            video_description_unit=unit_parsed,
            personal_tags=personal_tags,
            style_vector=style_vector_parsed,
            style_signature=style_signature_parsed,
            raw_outputs=raw_outputs,
        )
    except Exception as exc:
        code = video_tagging_error_code(exc)
        task = update_video_tagging_task(
            task_id,
            status="failed",
            result_code=code,
            result_message=str(exc),
            error_code=str(code),
            error_message=str(exc),
            raw_outputs=raw_outputs,
        )
    if task:
        send_video_tagging_callback(task)


def send_video_tagging_callback(task):
    payload = build_video_tagging_callback_payload(task)
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    attempts = int(task.get("callback_attempts") or 0) + 1
    try:
        request = urllib.request.Request(
            task["callback_url"],
            data=data,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=15) as response:
            body = response.read().decode("utf-8", errors="ignore")
            update_video_tagging_task(
                task["id"],
                callback_status="success" if 200 <= response.status < 300 else "failed",
                callback_response_code=response.status,
                callback_response_body=body[:2000],
                callback_attempts=attempts,
            )
    except Exception as exc:
        update_video_tagging_task(
            task["id"],
            callback_status="failed",
            callback_response_code=0,
            callback_response_body=str(exc)[:2000],
            callback_attempts=attempts,
        )


def video_queue_worker(worker_index):
    worker_id = f"video-{os.getpid()}-{worker_index}-{uuid.uuid4().hex[:8]}"
    poll_seconds = current_queue_poll_seconds()
    while True:
        try:
            task = claim_next_video_tagging_task(worker_id)
            if task:
                run_video_tagging_task(task["id"])
            else:
                time.sleep(poll_seconds)
        except Exception as exc:
            print(f"Video queue worker {worker_id} failed: {exc}")
            time.sleep(poll_seconds)


def blogger_queue_worker(worker_index):
    worker_id = f"blogger-{os.getpid()}-{worker_index}-{uuid.uuid4().hex[:8]}"
    poll_seconds = current_queue_poll_seconds()
    while True:
        try:
            task = claim_next_blogger_tagging_task(worker_id)
            if task:
                run_blogger_tagging_task(task["id"])
            else:
                time.sleep(poll_seconds)
        except Exception as exc:
            print(f"Blogger queue worker {worker_id} failed: {exc}")
            time.sleep(poll_seconds)


def start_queue_workers():
    global QUEUE_WORKERS_STARTED
    with QUEUE_WORKERS_LOCK:
        if QUEUE_WORKERS_STARTED:
            return
        QUEUE_WORKERS_STARTED = True
        for index in range(current_video_worker_count()):
            threading.Thread(target=video_queue_worker, args=(index + 1,), daemon=True).start()
        for index in range(current_blogger_worker_count()):
            threading.Thread(target=blogger_queue_worker, args=(index + 1,), daemon=True).start()


def load_video_rows():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Excel not found: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    required = ["TikTok 博主", "博主主页 URL", "视频原始链接"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"Excel missing columns: {', '.join(missing)}")

    index = {name: headers.index(name) for name in headers if name}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        blogger = row[index["TikTok 博主"]]
        profile_url = row[index["博主主页 URL"]]
        source_url = row[index["视频原始链接"]]
        if not blogger or not profile_url or not source_url:
            continue
        gcs_url = row[index["local_gcs_video_url"]] if "local_gcs_video_url" in index else ""
        caption = row[index["caption"]] if "caption" in index else ""
        hashtag = row[index["hashtag"]] if "hashtag" in index else ""
        file_uri = gcs_url or source_url
        if not file_uri:
            continue
        rows.append(
            {
                "blogger": str(blogger),
                "profile_url": str(profile_url),
                "source_url": str(source_url),
                "file_uri": str(file_uri),
                "caption": str(caption or ""),
                "hashtag": str(hashtag or ""),
            }
        )
    wb.close()
    return rows


def sample_bloggers(blogger_count=30, videos_per_blogger=15):
    grouped = {}
    for row in load_video_rows():
        key = row["profile_url"]
        grouped.setdefault(
            key, {"blogger": row["blogger"], "profile_url": row["profile_url"], "videos": []}
        )
        grouped[key]["videos"].append(row)

    candidates = [item for item in grouped.values() if len(item["videos"]) >= videos_per_blogger]
    random.shuffle(candidates)
    selected = candidates[:blogger_count]

    for idx, blogger in enumerate(selected, start=1):
        random.shuffle(blogger["videos"])
        blogger["videos"] = blogger["videos"][:videos_per_blogger]
        blogger["id"] = f"blogger_{idx}"
        blogger["status"] = "pending"
        blogger["progress"] = {"done": 0, "total": videos_per_blogger * 4 + 1}
        blogger["video_units"] = [None] * videos_per_blogger
        blogger["video_classifications"] = [None] * videos_per_blogger
        blogger["video_style_vectors"] = [None] * videos_per_blogger
        blogger["video_style_signatures"] = [None] * videos_per_blogger
        blogger["classification_summary"] = None
        blogger["account_style_summary"] = None
        blogger["account_result"] = None
        blogger["final_account_tags"] = None
        blogger["error"] = ""
    return selected


def normalize_blogger_results(bloggers):
    for blogger in bloggers:
        video_count = len(blogger.get("videos") or [])
        blogger.setdefault("video_units", [None] * video_count)
        blogger.setdefault("video_classifications", [None] * video_count)
        blogger.setdefault("video_style_vectors", [None] * video_count)
        blogger.setdefault("video_style_signatures", [None] * video_count)
        blogger.setdefault("classification_summary", None)
        blogger.setdefault("account_style_summary", None)
        blogger.setdefault("account_result", None)
        blogger.setdefault("final_account_tags", None)
        blogger.setdefault("error", "")
        blogger.setdefault("status", "pending")
        blogger["progress"] = blogger.get("progress") or {"done": 0, "total": video_count * 4 + 1}
        blogger["progress"]["total"] = video_count * 4 + 1
    return bloggers


def reset_blogger_results(bloggers):
    reset = json.loads(json.dumps(bloggers, ensure_ascii=False))
    for blogger in reset:
        video_count = len(blogger.get("videos") or [])
        blogger["status"] = "pending"
        blogger["progress"] = {"done": 0, "total": video_count * 4 + 1}
        blogger["video_units"] = [None] * video_count
        blogger["video_classifications"] = [None] * video_count
        blogger["video_style_vectors"] = [None] * video_count
        blogger["video_style_signatures"] = [None] * video_count
        blogger["classification_summary"] = None
        blogger["account_style_summary"] = None
        blogger["account_result"] = None
        blogger["final_account_tags"] = None
        blogger["error"] = ""
    return reset


def load_current_batch():
    if not CURRENT_BATCH_PATH.exists():
        return None
    return normalize_blogger_results(json.loads(CURRENT_BATCH_PATH.read_text()))


def save_current_batch(bloggers):
    CURRENT_BATCH_PATH.write_text(json.dumps(bloggers, ensure_ascii=False, indent=2))


def get_or_create_batch(blogger_count=30, videos_per_blogger=15, force=False):
    with BATCH_LOCK:
        if not force:
            existing = load_current_batch()
            if existing:
                return existing
        bloggers = sample_bloggers(blogger_count, videos_per_blogger)
        save_current_batch(bloggers)
        return bloggers


def mark_saved_job_inactive(job):
    if job.get("status") in {"pending", "running"}:
        job = json.loads(json.dumps(job, ensure_ascii=False))
        job["status"] = "interrupted"
        job["error"] = "服务已重启，旧任务后台线程已停止。请重新点击一键运行。"
        for blogger in job.get("bloggers", []):
            if blogger.get("status") == "running":
                blogger["status"] = "interrupted"
    return job


def extract_text(api_result):
    choices = api_result.get("choices", [])
    if choices:
        content = choices[0].get("message", {}).get("content", "")
        if isinstance(content, str):
            return content
        if isinstance(content, list):
            texts = []
            for part in content:
                if isinstance(part, dict):
                    texts.append(part.get("text") or part.get("content") or "")
            return "\n".join(text for text in texts if text)
    for candidate in api_result.get("candidates", []):
        for part in candidate.get("content", {}).get("parts", []):
            if "text" in part:
                return part["text"]
    return json.dumps(api_result, ensure_ascii=False)


def parse_json_text(text):
    if not text:
        return None
    cleaned = text.strip()
    fenced = re.search(r"```(?:json)?\s*(.*?)```", cleaned, re.S)
    if fenced:
        cleaned = fenced.group(1).strip()
    try:
        return json.loads(cleaned)
    except Exception:
        pass
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start >= 0 and end > start:
        try:
            return json.loads(cleaned[start : end + 1])
        except Exception:
            return None
    return None


def is_gcs_url(url):
    if not url:
        return False
    if str(url).startswith("gs://"):
        return True
    parsed = urlparse(str(url))
    return parsed.scheme in ("http", "https") and parsed.netloc in GCS_HTTPS_HOSTS


def parse_gcs_url(url):
    raw = str(url or "")
    if raw.startswith("gs://"):
        parsed = urlparse(raw)
        bucket = parsed.netloc
        key = parsed.path.lstrip("/")
    else:
        parsed = urlparse(raw)
        if parsed.netloc not in GCS_HTTPS_HOSTS:
            raise ValueError(f"not a GCS URL: {raw}")
        parts = parsed.path.lstrip("/").split("/", 1)
        if len(parts) != 2:
            raise ValueError(f"invalid GCS URL: {raw}")
        bucket, key = parts[0], unquote(parts[1])
    if not bucket or not key:
        raise ValueError(f"invalid GCS URL: {raw}")
    return bucket, key


def parse_gcs_signed_url_expiration(url):
    try:
        params = parse_qs(urlparse(str(url)).query)
        date_str = (params.get("X-Goog-Date") or [None])[0]
        expires_str = (params.get("X-Goog-Expires") or [None])[0]
        if not date_str or not expires_str:
            return None
        signed_at = datetime.strptime(date_str, "%Y%m%dT%H%M%SZ").replace(tzinfo=timezone.utc)
        return signed_at + timedelta(seconds=int(expires_str))
    except Exception:
        return None


def is_gcs_signed_url_fresh(url):
    expires_at = parse_gcs_signed_url_expiration(url)
    if expires_at is None:
        return False
    return expires_at > datetime.now(timezone.utc) + GCS_SIGNED_URL_REFRESH_THRESHOLD


def generate_gcs_signed_url(bucket_name, object_key, expiration_minutes=None):
    from google.cloud import storage

    expiration_minutes = expiration_minutes or int(GCS_SIGNED_URL_TTL.total_seconds() // 60)
    client = storage.Client()
    blob = client.bucket(bucket_name).blob(object_key)
    kwargs = {
        "version": "v4",
        "expiration": timedelta(minutes=expiration_minutes),
        "method": "GET",
    }
    try:
        return blob.generate_signed_url(**kwargs)
    except AttributeError as exc:
        if "private key" not in str(exc).lower():
            raise

    import google.auth
    from google.auth.transport import requests as auth_requests

    credentials, _ = google.auth.default()
    credentials.refresh(auth_requests.Request())
    service_account_email = getattr(credentials, "service_account_email", None)
    if not service_account_email:
        raise RuntimeError("ADC credentials do not include service_account_email for GCS signing")
    return blob.generate_signed_url(
        **kwargs,
        service_account_email=service_account_email,
        access_token=credentials.token,
    )


def ensure_fresh_gcs_signed_url(url):
    if not is_gcs_url(url):
        return url
    if is_gcs_signed_url_fresh(url):
        return url
    bucket, key = parse_gcs_url(url)
    return generate_gcs_signed_url(bucket, key)


def call_evolink_text(messages, timeout=180):
    api_key = current_api_key()
    if not api_key:
        raise RuntimeError("Missing EVOLINK_API_KEY")
    payload = {"model": current_text_model(), "messages": messages}
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        current_text_api_url(),
        data=data,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"API HTTP {exc.code}: {detail}") from exc


def call_evolink_video(file_uri, prompt, timeout=180):
    api_key = current_api_key()
    if not api_key:
        raise RuntimeError("Missing EVOLINK_API_KEY")
    file_uri = ensure_fresh_gcs_signed_url(file_uri)
    api_url = current_video_api_url()
    if "/chat/completions" in api_url:
        payload = {
            "model": current_text_model(),
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": file_uri}},
                    ],
                }
            ],
        }
    else:
        payload = {
            "contents": [
                {
                    "role": "user",
                    "parts": [
                        {"fileData": {"mimeType": "video/mp4", "fileUri": file_uri}},
                        {"text": prompt},
                    ],
                }
            ]
        }
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        api_url,
        data=data,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": "curl/8.7.1",
            "Accept": "*/*",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"Video API HTTP {exc.code}: {detail}") from exc


def build_video_prompt(base_prompt, video, video_index):
    return (
        f"{base_prompt}\n\n"
        f"本条视频 metadata：\n"
        f"- video_index: {video_index}\n"
        f"- video_url: {video['source_url']}\n"
        f"- caption: {video['caption'] or '无'}\n"
        f"- hashtags: {video['hashtag'] or '无'}\n\n"
        f"请把以上 caption 和 hashtags 写入 social_media_info 字段。"
    )


def build_account_prompt(base_prompt, blogger):
    units = [
        unit.get("parsed") or unit.get("raw_text")
        for unit in blogger["video_units"]
        if unit is not None and not unit.get("error")
    ]
    return (
        f"{base_prompt}\n\n"
        f"账号信息：\n"
        f"- blogger: {blogger['blogger']}\n"
        f"- profile_url: {blogger['profile_url']}\n\n"
        f"以下是该账号的 {len(units)} 个 video_description_unit：\n"
        f"{json.dumps(units, ensure_ascii=False, indent=2)}"
    )


def build_classification_prompt(base_prompt, video, unit):
    unit_payload = unit.get("parsed") or unit.get("raw_text") or {}
    return (
        f"{base_prompt}\n\n"
        f"本条视频 metadata：\n"
        f"- video_url: {video['source_url']}\n"
        f"- caption: {video['caption'] or '无'}\n"
        f"- hashtags: {video['hashtag'] or '无'}\n\n"
        f"第一阶段 video_description_unit：\n"
        f"{json.dumps(unit_payload, ensure_ascii=False, indent=2)}"
    )


def build_style_vector_prompt(base_prompt, video, unit):
    unit_payload = unit.get("parsed") or unit.get("raw_text") or {}
    return (
        f"{base_prompt}\n\n"
        f"本条视频 metadata：\n"
        f"- video_url: {video['source_url']}\n"
        f"- caption: {video['caption'] or '无'}\n"
        f"- hashtags: {video['hashtag'] or '无'}\n\n"
        f"第一阶段 video_description_unit：\n"
        f"{json.dumps(unit_payload, ensure_ascii=False, indent=2)}"
    )


def build_style_signature_prompt(base_prompt, video, unit, style_vector=None):
    unit_payload = unit.get("parsed") or unit.get("raw_text") or {}
    vector_payload = style_vector.get("parsed") if isinstance(style_vector, dict) else None
    return (
        f"{base_prompt}\n\n"
        f"本条视频 metadata：\n"
        f"- video_url: {video['source_url']}\n"
        f"- caption: {video['caption'] or '无'}\n"
        f"- hashtags: {video['hashtag'] or '无'}\n\n"
        f"第一阶段 video_description_unit：\n"
        f"{json.dumps(unit_payload, ensure_ascii=False, indent=2)}\n\n"
        f"可选 style_vector：\n"
        f"{json.dumps(vector_payload or {}, ensure_ascii=False, indent=2)}"
    )


def update_job(job_id, mutator):
    with JOBS_LOCK:
        job = JOBS[job_id]
        mutator(job)
        snapshot = json.loads(json.dumps(job, ensure_ascii=False))
    (DATA_DIR / f"{job_id}.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2))
    CURRENT_JOB_PATH.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2))


def analyze_video_unit(prompt1, video, video_index, api_gate):
    base = {
        "video_index": video_index,
        "video_url": video["source_url"],
        "file_uri": video["file_uri"],
        "caption": video["caption"],
        "hashtag": video["hashtag"],
    }
    try:
        with api_gate:
            result = call_evolink_video(
                video["file_uri"],
                build_video_prompt(prompt1, video, video_index),
            )
        text = extract_text(result)
        return {
            **base,
            "raw_text": text,
            "parsed": parse_json_text(text),
            "error": "",
        }
    except Exception as exc:
        return {
            **base,
            "raw_text": "",
            "parsed": None,
            "error": str(exc),
        }


def analyze_account(prompt2, blogger, api_gate):
    valid_units = [unit for unit in blogger["video_units"] if unit is not None and not unit.get("error")]
    if not valid_units:
        return {
            "raw_text": "",
            "parsed": None,
            "error": "No successful video units for this blogger.",
    }
    try:
        messages = [{"role": "user", "content": build_account_prompt(prompt2, blogger)}]
        with api_gate:
            account_result = call_evolink_text(messages)
        account_text = extract_text(account_result)
        return {
            "raw_text": account_text,
            "parsed": parse_json_text(account_text),
            "error": "",
        }
    except Exception as exc:
        return {
            "raw_text": "",
            "parsed": None,
            "error": str(exc),
        }


def analyze_video_classification(prompt3, video, unit, api_gate):
    base = {
        "video_index": unit.get("video_index"),
        "video_url": video["source_url"],
        "caption": video["caption"],
        "hashtag": video["hashtag"],
    }
    if unit.get("error"):
        return {**base, "raw_text": "", "parsed": None, "error": unit["error"]}
    try:
        messages = [{"role": "user", "content": build_classification_prompt(prompt3, video, unit)}]
        with api_gate:
            result = call_evolink_text(messages)
        text = extract_text(result)
        return {**base, "raw_text": text, "parsed": parse_json_text(text), "error": ""}
    except Exception as exc:
        return {**base, "raw_text": "", "parsed": None, "error": str(exc)}


def analyze_video_style_vector(prompt4, video, unit, api_gate):
    base = {
        "video_index": unit.get("video_index"),
        "video_url": video["source_url"],
        "caption": video["caption"],
        "hashtag": video["hashtag"],
    }
    if unit.get("error"):
        return {**base, "raw_text": "", "parsed": None, "error": unit["error"]}
    try:
        messages = [{"role": "user", "content": build_style_vector_prompt(prompt4, video, unit)}]
        with api_gate:
            result = call_evolink_text(messages)
        text = extract_text(result)
        return {**base, "raw_text": text, "parsed": parse_json_text(text), "error": ""}
    except Exception as exc:
        return {**base, "raw_text": "", "parsed": None, "error": str(exc)}


def analyze_video_style_signature(prompt5, video, unit, style_vector, api_gate):
    base = {
        "video_index": unit.get("video_index"),
        "video_url": video["source_url"],
        "caption": video["caption"],
        "hashtag": video["hashtag"],
    }
    if unit.get("error"):
        return {**base, "raw_text": "", "parsed": None, "error": unit["error"]}
    try:
        messages = [
            {
                "role": "user",
                "content": build_style_signature_prompt(prompt5, video, unit, style_vector),
            }
        ]
        with api_gate:
            result = call_evolink_text(messages)
        text = extract_text(result)
        return {**base, "raw_text": text, "parsed": parse_json_text(text), "error": ""}
    except Exception as exc:
        return {**base, "raw_text": "", "parsed": None, "error": str(exc)}


def extract_classification_label(classification, key):
    parsed = classification.get("parsed") if classification else None
    if not isinstance(parsed, dict):
        return None
    block = parsed.get(key) or {}
    label = block.get("label")
    return label if isinstance(label, str) and label else None


def distribution_for(classifications, key, remix_label, unclear_label):
    labels = [extract_classification_label(item, key) for item in classifications if item and not item.get("error")]
    labels = [label for label in labels if label]
    total = len(labels)
    counts = {}
    for label in labels:
        counts[label] = counts.get(label, 0) + 1
    rows = [
        {
            "label": label,
            "count": count,
            "share": f"{round(count * 100 / total)}%" if total else "0%",
        }
        for label, count in sorted(counts.items(), key=lambda item: item[1], reverse=True)
    ]
    if not rows:
        final_label = unclear_label
    else:
        top = rows[0]["count"] / total
        second = rows[1]["count"] / total if len(rows) > 1 else 0
        final_label = rows[0]["label"] if top >= 0.5 and second <= 0.4 else remix_label
    return {"final_label": final_label, "total": total, "distribution": rows}


def aggregate_classifications(classifications):
    return {
        "social_identity": distribution_for(
            classifications,
            "social_identity_classification",
            "Remix 身份型",
            "无明确社会身份型",
        ),
        "occasion": distribution_for(
            classifications,
            "occasion_classification",
            "Remix Occasion",
            "无明显Occasion",
        ),
    }


def final_account_tags(account_result, classification_summary):
    parsed = account_result.get("parsed") if account_result else None
    if not isinstance(parsed, dict):
        return None
    tags = parsed.get("account_level_personal_tags")
    if not isinstance(tags, dict):
        return None
    merged = json.loads(json.dumps(tags, ensure_ascii=False))
    social_identity = (classification_summary or {}).get("social_identity", {}).get("final_label")
    occasion = (classification_summary or {}).get("occasion", {}).get("final_label")
    if social_identity:
        merged["social_identity"] = social_identity
    if occasion:
        merged["occasion"] = occasion
    merged.pop("aesthetic_style", None)
    return merged


def refresh_derived_fields(job):
    for blogger in job.get("bloggers", []):
        if blogger.get("video_style_vectors") or blogger.get("video_style_signatures"):
            blogger["account_style_summary"] = aggregate_style_results(
                blogger.get("video_style_vectors") or [],
                blogger.get("video_style_signatures") or [],
            )
        if blogger.get("account_result"):
            blogger["final_account_tags"] = final_account_tags(
                blogger.get("account_result"),
                blogger.get("classification_summary"),
            )
    return job


def parsed_payload(result, wrapper_key=None):
    parsed = result.get("parsed") if isinstance(result, dict) else None
    if not isinstance(parsed, dict):
        return {}
    if wrapper_key and isinstance(parsed.get(wrapper_key), dict):
        return parsed[wrapper_key]
    return parsed


def top_counts(values, limit=5):
    counts = {}
    first_seen = {}
    for value in values:
        if value in ("", None, "null"):
            continue
        if not isinstance(value, (str, int, float, bool)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        if value not in first_seen:
            first_seen[value] = len(first_seen)
        counts[value] = counts.get(value, 0) + 1
    return [
        value
        for value, _ in sorted(counts.items(), key=lambda item: (-item[1], first_seen[item[0]]))[:limit]
    ]


def mode_value(values, default=""):
    top = top_counts(values, limit=1)
    return top[0] if top else default


def numeric_mean(values, default=0.0):
    nums = [float(value) for value in values if isinstance(value, (int, float))]
    if not nums:
        return default
    return round(sum(nums) / len(nums), 4)


def frequent_values(values, total, limit=5, min_count=2, min_share=0.3):
    counts = {}
    first_seen = {}
    for value in values:
        if value in ("", None, "null"):
            continue
        if not isinstance(value, (str, int, float, bool)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        if value not in first_seen:
            first_seen[value] = len(first_seen)
        counts[value] = counts.get(value, 0) + 1
    threshold = 1 if total < 3 else max(min_count, int(total * min_share + 0.9999))
    rows = sorted(counts.items(), key=lambda item: (-item[1], first_seen[item[0]]))
    filtered = [value for value, count in rows if count >= threshold]
    if not filtered and rows:
        filtered = [rows[0][0]]
    return filtered[:limit]


def majority_bool(values):
    bools = [value for value in values if isinstance(value, bool)]
    if not bools:
        return False
    return sum(1 for value in bools if value) / len(bools) >= 0.5


def mixed_or_mode(values, default="mixed"):
    filtered = [value for value in values if value not in ("", None, "null")]
    if not filtered:
        return default
    counts = {}
    first_seen = {}
    for value in filtered:
        if value not in first_seen:
            first_seen[value] = len(first_seen)
        counts[value] = counts.get(value, 0) + 1
    top_value, top_count = sorted(counts.items(), key=lambda item: (-item[1], first_seen[item[0]]))[0]
    return top_value if top_count / len(filtered) >= 0.5 else "mixed"


def aggregate_length_preferences(values):
    buckets = {}
    for item in values:
        if not isinstance(item, dict):
            continue
        for key, value in item.items():
            if value in ("", None, "null"):
                continue
            buckets.setdefault(key, []).append(value)
    return {key: mode_value(items) for key, items in buckets.items()}


def collect_style_signatures(style_signatures):
    signatures = []
    for item in style_signatures or []:
        if not item or item.get("error"):
            continue
        sig = parsed_payload(item, "style_signature")
        if sig:
            signatures.append(sig)
    return signatures


def aggregate_account_style_signature(signatures):
    total = len(signatures)
    color = [sig.get("color_palette") or {} for sig in signatures]
    material = [sig.get("material_profile") or {} for sig in signatures]
    silhouette = [sig.get("silhouette_profile") or {} for sig in signatures]
    pattern = [sig.get("pattern_profile") or {} for sig in signatures]
    mood = [sig.get("aesthetic_mood") or {} for sig in signatures]
    price = [sig.get("price_positioning") or {} for sig in signatures]
    era = [sig.get("era_influence") or {} for sig in signatures]

    occasion_totals = {}
    occasion_count = 0
    for sig in signatures:
        occasion = sig.get("occasion_vector") or {}
        if not isinstance(occasion, dict):
            continue
        occasion_count += 1
        for key, value in occasion.items():
            if isinstance(value, (int, float)):
                occasion_totals[key] = occasion_totals.get(key, 0.0) + float(value)
    occasion_vector = {
        key: round(value / occasion_count, 4)
        for key, value in sorted(occasion_totals.items())
    } if occasion_count else {}

    primary_era = mode_value(
        [item.get("primary_era") for item in era if item.get("primary_era") is not None],
        None,
    )
    era_authenticity = mode_value(
        [item.get("era_authenticity") for item in era if item.get("era_authenticity") is not None],
        None,
    )

    return {
        "color_palette": {
            "dominant_colors": frequent_values(
                [value for item in color for value in (item.get("dominant_colors") or [])],
                total,
            ),
            "temperature": mode_value([item.get("temperature") for item in color], "neutral"),
            "saturation": mode_value([item.get("saturation") for item in color], "muted"),
            "contrast": mode_value([item.get("contrast") for item in color], "medium"),
            "signature_combos": frequent_values(
                [value for item in color for value in (item.get("signature_combos") or [])],
                total,
                limit=3,
            ),
            "monochromatic_tendency": numeric_mean(
                [item.get("monochromatic_tendency") for item in color]
            ),
        },
        "material_profile": {
            "primary_materials": frequent_values(
                [value for item in material for value in (item.get("primary_materials") or [])],
                total,
            ),
            "texture_preference": mode_value(
                [item.get("texture_preference") for item in material],
                "mixed",
            ),
            "weight_preference": mode_value(
                [item.get("weight_preference") for item in material],
                "medium",
            ),
            "transparency_level": mode_value(
                [item.get("transparency_level") for item in material],
                "opaque",
            ),
            "hardware_affinity": numeric_mean(
                [item.get("hardware_affinity") for item in material]
            ),
        },
        "silhouette_profile": {
            "fit_preference": mode_value(
                [item.get("fit_preference") for item in silhouette],
                "regular",
            ),
            "proportion_play": mode_value(
                [item.get("proportion_play") for item in silhouette],
                "balanced",
            ),
            "structure_level": mode_value(
                [item.get("structure_level") for item in silhouette],
                "semi-structured",
            ),
            "length_preference": aggregate_length_preferences(
                [item.get("length_preference") for item in silhouette]
            ),
            "layering_complexity": mode_value(
                [item.get("layering_complexity") for item in silhouette],
                "minimal",
            ),
        },
        "pattern_profile": {
            "pattern_types": frequent_values(
                [value for item in pattern for value in (item.get("pattern_types") or [])],
                total,
            ),
            "pattern_scale": mode_value(
                [item.get("pattern_scale") for item in pattern],
                "medium",
            ),
            "pattern_frequency": numeric_mean(
                [item.get("pattern_frequency") for item in pattern]
            ),
            "logo_visibility": mode_value(
                [item.get("logo_visibility") for item in pattern],
                "none",
            ),
            "print_mixing": majority_bool(
                [item.get("print_mixing") for item in pattern]
            ),
        },
        "aesthetic_mood": {
            "energy": mode_value([item.get("energy") for item in mood], "balanced"),
            "formality_range": [
                mode_value(
                    [item.get("formality_range", [None, None])[0] for item in mood if isinstance(item.get("formality_range"), list)],
                    "casual",
                ),
                mode_value(
                    [item.get("formality_range", [None, None])[1] for item in mood if isinstance(item.get("formality_range"), list) and len(item.get("formality_range")) > 1],
                    "semi-formal",
                ),
            ],
            "gender_expression": mode_value(
                [item.get("gender_expression") for item in mood],
                "fluid",
            ),
            "cultural_references": frequent_values(
                [value for item in mood for value in (item.get("cultural_references") or [])],
                total,
            ),
            "mood_keywords": frequent_values(
                [value for item in mood for value in (item.get("mood_keywords") or [])],
                total,
            ),
        },
        "occasion_vector": occasion_vector,
        "price_positioning": {
            "tier": mixed_or_mode([item.get("tier") for item in price], "mid-range"),
            "investment_vs_trend": numeric_mean(
                [item.get("investment_vs_trend") for item in price]
            ),
            "brand_consciousness": numeric_mean(
                [item.get("brand_consciousness") for item in price]
            ),
        },
        "era_influence": {
            "primary_era": primary_era,
            "era_authenticity": era_authenticity,
            "retro_futurism": numeric_mean(
                [item.get("retro_futurism") for item in era]
            ),
        },
    }


def aggregate_style_results(style_vectors, style_signatures):
    vector_totals = {}
    vector_count = 0
    for item in style_vectors or []:
        if not item or item.get("error"):
            continue
        vector = parsed_payload(item, "style_vector")
        numeric = {
            key: float(value)
            for key, value in vector.items()
            if key in STYLE_NAMES and isinstance(value, (int, float))
        }
        if not numeric:
            continue
        vector_count += 1
        for style in STYLE_NAMES:
            vector_totals[style] = vector_totals.get(style, 0.0) + numeric.get(style, 0.0)

    average_vector = {
        style: round(vector_totals.get(style, 0.0) / vector_count, 4)
        for style in STYLE_NAMES
    } if vector_count else {}
    top_styles = [
        {"style": style, "score": score}
        for style, score in sorted(average_vector.items(), key=lambda item: item[1], reverse=True)[:5]
        if score > 0
    ]

    signatures = collect_style_signatures(style_signatures)
    account_style_signature = aggregate_account_style_signature(signatures)
    colors = []
    materials = []
    moods = []
    fit_values = []
    price_tiers = []
    eras = []
    temperatures = []
    occasion_totals = {}
    occasion_count = 0

    for sig in signatures:
        colors.extend(sig.get("color_palette", {}).get("dominant_colors") or [])
        temperatures.append(sig.get("color_palette", {}).get("temperature"))
        materials.extend(sig.get("material_profile", {}).get("primary_materials") or [])
        fit_values.append(sig.get("silhouette_profile", {}).get("fit_preference"))
        moods.extend(sig.get("aesthetic_mood", {}).get("mood_keywords") or [])
        price_tiers.append(sig.get("price_positioning", {}).get("tier"))
        eras.append(sig.get("era_influence", {}).get("primary_era"))
        occasion = sig.get("occasion_vector") or {}
        if isinstance(occasion, dict):
            occasion_count += 1
            for key, value in occasion.items():
                if isinstance(value, (int, float)):
                    occasion_totals[key] = occasion_totals.get(key, 0.0) + float(value)

    top_occasions = []
    if occasion_count:
        top_occasions = [
            {"occasion": key, "score": round(value / occasion_count, 4)}
            for key, value in sorted(occasion_totals.items(), key=lambda item: item[1], reverse=True)[:5]
        ]

    return {
        "video_count": max(vector_count, len(signatures)),
        "average_style_vector": average_vector,
        "top_styles": top_styles,
        "account_style_signature": account_style_signature,
        "dominant_colors": top_counts(colors),
        "temperature": mode_value(temperatures),
        "primary_materials": top_counts(materials),
        "fit_preference": mode_value(fit_values),
        "mood_keywords": top_counts(moods),
        "top_occasions": top_occasions,
        "price_tier": mode_value(price_tiers),
        "primary_era": mode_value(eras),
    }


def run_job(job_id, prompt1, prompt2, prompt3, prompt4, prompt5, concurrency, api_concurrency):
    try:
        update_job(job_id, lambda job: job.update({"status": "running", "started_at": time.time()}))
        with JOBS_LOCK:
            bloggers = JOBS[job_id]["bloggers"]
        api_gate = threading.BoundedSemaphore(api_concurrency)

        def mark_all_running(job):
            job["concurrency"] = concurrency
            job["api_concurrency"] = api_concurrency
            for blogger in job["bloggers"]:
                blogger["status"] = "running"

        update_job(job_id, mark_all_running)

        with ThreadPoolExecutor(max_workers=concurrency) as video_executor, ThreadPoolExecutor(
            max_workers=concurrency
        ) as text_executor:
            video_futures = {}
            classification_futures = {}
            style_vector_futures = {}
            style_signature_futures = {}
            account_futures = {}
            account_submitted = set()

            def update_blogger_completion(job, blogger_index):
                b = job["bloggers"][blogger_index]
                video_done = sum(item is not None for item in b["video_units"])
                class_done = sum(item is not None for item in b["video_classifications"])
                vector_done = sum(item is not None for item in b["video_style_vectors"])
                signature_done = sum(item is not None for item in b["video_style_signatures"])
                account_done = b.get("account_result") is not None
                b["progress"]["total"] = len(b.get("videos") or []) * 4 + 1
                b["progress"]["done"] = video_done + class_done + vector_done + signature_done + (1 if account_done else 0)
                if account_done:
                    b["final_account_tags"] = final_account_tags(
                        b.get("account_result"),
                        b.get("classification_summary"),
                    )
                if vector_done or signature_done:
                    b["account_style_summary"] = aggregate_style_results(
                        b.get("video_style_vectors") or [],
                        b.get("video_style_signatures") or [],
                    )
                expected = len(b.get("videos") or [])
                if account_done and class_done == expected and vector_done == expected and signature_done == expected:
                    b["status"] = "done" if not b["account_result"].get("error") else "partial"
                else:
                    b["status"] = "running"

            for blogger_index, blogger in enumerate(bloggers):
                for video_index, video in enumerate(blogger["videos"], start=1):
                    future = video_executor.submit(analyze_video_unit, prompt1, video, video_index, api_gate)
                    video_futures[future] = (blogger_index, video_index)

            while video_futures or classification_futures or style_vector_futures or style_signature_futures or account_futures:
                active_futures = (
                    list(video_futures.keys())
                    + list(classification_futures.keys())
                    + list(style_vector_futures.keys())
                    + list(style_signature_futures.keys())
                    + list(account_futures.keys())
                )
                done, _ = wait(
                    active_futures,
                    return_when=FIRST_COMPLETED,
                )
                for future in done:
                    if future in video_futures:
                        blogger_index, video_index = video_futures.pop(future)
                        try:
                            unit = future.result()
                        except Exception as exc:
                            unit = {
                                "video_index": video_index,
                                "video_url": bloggers[blogger_index]["videos"][video_index - 1]["source_url"],
                                "file_uri": bloggers[blogger_index]["videos"][video_index - 1]["file_uri"],
                                "caption": bloggers[blogger_index]["videos"][video_index - 1]["caption"],
                                "hashtag": bloggers[blogger_index]["videos"][video_index - 1]["hashtag"],
                                "raw_text": "",
                                "parsed": None,
                                "error": str(exc),
                            }

                        def save_unit(job):
                            b = job["bloggers"][blogger_index]
                            b["video_units"][video_index - 1] = unit
                            update_blogger_completion(job, blogger_index)

                        update_job(job_id, save_unit)

                        with JOBS_LOCK:
                            units_done = all(
                                item is not None for item in JOBS[job_id]["bloggers"][blogger_index]["video_units"]
                            )
                            if units_done and blogger_index not in account_submitted:
                                account_submitted.add(blogger_index)
                                blogger_snapshot = json.loads(
                                    json.dumps(JOBS[job_id]["bloggers"][blogger_index], ensure_ascii=False)
                                )
                            else:
                                blogger_snapshot = None

                        if blogger_snapshot is not None:
                            account_future = text_executor.submit(analyze_account, prompt2, blogger_snapshot, api_gate)
                            account_futures[account_future] = blogger_index

                        class_future = text_executor.submit(
                            analyze_video_classification,
                            prompt3,
                            bloggers[blogger_index]["videos"][video_index - 1],
                            unit,
                            api_gate,
                        )
                        classification_futures[class_future] = (blogger_index, video_index)

                        style_future = text_executor.submit(
                            analyze_video_style_vector,
                            prompt4,
                            bloggers[blogger_index]["videos"][video_index - 1],
                            unit,
                            api_gate,
                        )
                        style_vector_futures[style_future] = (blogger_index, video_index, unit)

                    elif future in classification_futures:
                        blogger_index, video_index = classification_futures.pop(future)
                        try:
                            classification = future.result()
                        except Exception as exc:
                            classification = {
                                "video_index": video_index,
                                "video_url": bloggers[blogger_index]["videos"][video_index - 1]["source_url"],
                                "caption": bloggers[blogger_index]["videos"][video_index - 1]["caption"],
                                "hashtag": bloggers[blogger_index]["videos"][video_index - 1]["hashtag"],
                                "raw_text": "",
                                "parsed": None,
                                "error": str(exc),
                            }

                        def save_classification(job):
                            b = job["bloggers"][blogger_index]
                            b["video_classifications"][video_index - 1] = classification
                            b["classification_summary"] = aggregate_classifications(
                                b["video_classifications"]
                            )
                            update_blogger_completion(job, blogger_index)

                        update_job(job_id, save_classification)

                    elif future in style_vector_futures:
                        blogger_index, video_index, unit = style_vector_futures.pop(future)
                        try:
                            style_vector = future.result()
                        except Exception as exc:
                            style_vector = {
                                "video_index": video_index,
                                "video_url": bloggers[blogger_index]["videos"][video_index - 1]["source_url"],
                                "caption": bloggers[blogger_index]["videos"][video_index - 1]["caption"],
                                "hashtag": bloggers[blogger_index]["videos"][video_index - 1]["hashtag"],
                                "raw_text": "",
                                "parsed": None,
                                "error": str(exc),
                            }

                        def save_style_vector(job):
                            b = job["bloggers"][blogger_index]
                            b["video_style_vectors"][video_index - 1] = style_vector
                            b["account_style_summary"] = aggregate_style_results(
                                b["video_style_vectors"],
                                b["video_style_signatures"],
                            )
                            update_blogger_completion(job, blogger_index)

                        update_job(job_id, save_style_vector)

                        signature_future = text_executor.submit(
                            analyze_video_style_signature,
                            prompt5,
                            bloggers[blogger_index]["videos"][video_index - 1],
                            unit,
                            style_vector,
                            api_gate,
                        )
                        style_signature_futures[signature_future] = (blogger_index, video_index)

                    elif future in style_signature_futures:
                        blogger_index, video_index = style_signature_futures.pop(future)
                        try:
                            style_signature = future.result()
                        except Exception as exc:
                            style_signature = {
                                "video_index": video_index,
                                "video_url": bloggers[blogger_index]["videos"][video_index - 1]["source_url"],
                                "caption": bloggers[blogger_index]["videos"][video_index - 1]["caption"],
                                "hashtag": bloggers[blogger_index]["videos"][video_index - 1]["hashtag"],
                                "raw_text": "",
                                "parsed": None,
                                "error": str(exc),
                            }

                        def save_style_signature(job):
                            b = job["bloggers"][blogger_index]
                            b["video_style_signatures"][video_index - 1] = style_signature
                            b["account_style_summary"] = aggregate_style_results(
                                b["video_style_vectors"],
                                b["video_style_signatures"],
                            )
                            update_blogger_completion(job, blogger_index)

                        update_job(job_id, save_style_signature)

                    else:
                        blogger_index = account_futures.pop(future)
                        try:
                            account_result = future.result()
                        except Exception as exc:
                            account_result = {"raw_text": "", "parsed": None, "error": str(exc)}

                        def save_account(job):
                            b = job["bloggers"][blogger_index]
                            b["account_result"] = account_result
                            update_blogger_completion(job, blogger_index)

                        update_job(job_id, save_account)

        update_job(job_id, lambda job: job.update({"status": "done", "finished_at": time.time()}))
    except Exception as exc:
        detail = f"{exc}\n{traceback.format_exc()}"

        def fail(job):
            job["status"] = "failed"
            job["error"] = detail

        update_job(job_id, fail)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        return

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/api/config":
            json_response(self, public_config_payload())
            return
        if path == "/api/sample":
            params = parse_qs(parsed.query)
            count = int(params.get("count", ["30"])[0])
            videos = int(params.get("videos", ["15"])[0])
            force = params.get("force", ["0"])[0] == "1"
            try:
                bloggers = get_or_create_batch(count, videos, force=force)
                json_response(self, {"bloggers": bloggers, "available": len(bloggers)})
            except Exception as exc:
                error_response(self, str(exc), 500)
            return
        if path == "/api/current-job":
            if CURRENT_JOB_PATH.exists():
                json_response(
                    self,
                    refresh_derived_fields(
                        mark_saved_job_inactive(json.loads(CURRENT_JOB_PATH.read_text()))
                    ),
                )
                return
            error_response(self, "No current job", 404)
            return
        if path.startswith("/api/jobs/"):
            job_id = path.rsplit("/", 1)[-1]
            with JOBS_LOCK:
                job = JOBS.get(job_id)
            if not job:
                saved = DATA_DIR / f"{job_id}.json"
                if saved.exists():
                    job = mark_saved_job_inactive(json.loads(saved.read_text()))
            if not job:
                error_response(self, "Job not found", 404)
                return
            json_response(self, refresh_derived_fields(job))
            return
        if path == "/api/v1/video-tagging/tasks":
            params = parse_qs(parsed.query)
            status = (params.get("status", [""])[0] or "").strip() or None
            limit = int(params.get("limit", ["100"])[0])
            date = (params.get("date", [""])[0] or "").strip() or None
            try:
                tasks = list_video_tagging_tasks(status=status, limit=limit, date=date)
                api_code_response(self, 0, "success", data={"tasks": tasks})
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if path == "/api/v1/blogger-tagging/tasks":
            params = parse_qs(parsed.query)
            status = (params.get("status", [""])[0] or "").strip() or None
            limit = int(params.get("limit", ["100"])[0])
            date = (params.get("date", [""])[0] or "").strip() or None
            try:
                tasks = list_blogger_tagging_tasks(status=status, limit=limit, date=date)
                api_code_response(self, 0, "success", data={"tasks": tasks})
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        blogger_video_match = re.fullmatch(r"/api/v1/bloggers/([0-9a-fA-F-]+)/video-tagging/tasks", path)
        if blogger_video_match:
            tiktok_blogger_id = blogger_video_match.group(1)
            try:
                uuid.UUID(tiktok_blogger_id)
            except Exception:
                api_code_response(self, 4003, "tiktok_blogger_id must be a UUID", http_status=400)
                return
            params = parse_qs(parsed.query)
            status = (params.get("status", [""])[0] or "").strip() or None
            limit = int(params.get("limit", ["100"])[0])
            date = (params.get("date", [""])[0] or "").strip() or None
            try:
                tasks = list_blogger_video_tagging_tasks(
                    tiktok_blogger_id,
                    status=status,
                    limit=limit,
                    date=date,
                )
                api_code_response(self, 0, "success", data={"tasks": tasks})
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if path == "/api/v1/tagging/tasks":
            params = parse_qs(parsed.query)
            status = (params.get("status", [""])[0] or "").strip() or None
            task_type = (params.get("type", [""])[0] or "").strip()
            limit = int(params.get("limit", ["100"])[0])
            date = (params.get("date", [""])[0] or "").strip() or None
            try:
                video_tasks = [] if task_type == "blogger" else list_video_tagging_tasks(status=status, limit=limit, date=date)
                blogger_tasks = [] if task_type == "video" else list_blogger_tagging_tasks(status=status, limit=limit, date=date)
                api_code_response(
                    self,
                    0,
                    "success",
                    data={"video_tasks": video_tasks, "blogger_tasks": blogger_tasks},
                )
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if path.startswith("/api/v1/videos/signed-url/"):
            video_id = path.rsplit("/", 1)[-1]
            try:
                uuid.UUID(video_id)
            except Exception:
                api_code_response(self, 4005, "video_id must be a UUID", http_status=400)
                return
            try:
                data = fetch_video_signed_url(video_id)
                if not data:
                    api_code_response(self, 4041, "video tagging result not found", http_status=404)
                    return
                api_code_response(self, 0, "success", data=data)
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if path.startswith("/api/v1/videos/tag/"):
            video_id = path.rsplit("/", 1)[-1]
            try:
                uuid.UUID(video_id)
            except Exception:
                api_code_response(self, 4005, "video_id must be a UUID", http_status=400)
                return
            try:
                task = fetch_video_tagging_task(video_id)
                if not task:
                    api_code_response(self, 4041, "video tagging result not found", http_status=404)
                    return
                api_code_response(self, 0, "success", data=task)
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if path.startswith("/api/v1/bloggers/tag/"):
            tiktok_blogger_id = path.rsplit("/", 1)[-1]
            try:
                uuid.UUID(tiktok_blogger_id)
            except Exception:
                api_code_response(self, 4003, "tiktok_blogger_id must be a UUID", http_status=400)
                return
            try:
                task = fetch_blogger_tagging_task(tiktok_blogger_id)
                if not task:
                    api_code_response(self, 4042, "blogger tagging result not found", http_status=404)
                    return
                api_code_response(self, 0, "success", data=task)
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return

        static_path = STATIC_DIR / ("index.html" if path == "/" else path.lstrip("/"))
        if not static_path.exists() or not static_path.resolve().is_relative_to(STATIC_DIR):
            self.send_response(404)
            self.end_headers()
            return
        content_type = "text/html; charset=utf-8"
        if static_path.suffix == ".css":
            content_type = "text/css; charset=utf-8"
        elif static_path.suffix == ".js":
            content_type = "application/javascript; charset=utf-8"
        body = static_path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/config":
            try:
                body = read_json_body(self)
                json_response(self, {"code": 0, "message": "saved", "data": update_service_config(body)})
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if parsed.path == "/api/internal/tagging-callback":
            api_code_response(self, 0, "ok")
            return
        if parsed.path == "/api/v1/videos/tag":
            try:
                body = read_json_body(self)
                ok, error = validate_video_tagging_request(body)
                if not ok:
                    api_code_response(self, error["code"], error["message"], http_status=400)
                    return
                result = submit_video_tagging_task(
                    {
                        "video_id": str(body["video_id"]).strip(),
                        "gcs_url": str(body["gcs_url"]).strip(),
                        "description": str(body["description"]).strip(),
                        "callback_url": str(body["callback_url"]).strip(),
                    }
                )
                task = result["task"]
                api_code_response(
                    self,
                    result["code"],
                    result["message"],
                    task_id=task["id"],
                    video_id=task["video_id"],
                    status=task["status"],
                )
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return
        if parsed.path == "/api/v1/bloggers/tag":
            try:
                body = read_json_body(self)
                ok, error, data = validate_blogger_tagging_request(body)
                if not ok:
                    api_code_response(self, error["code"], error["message"], http_status=400)
                    return
                result = submit_blogger_tagging_task(data)
                if result.get("code") != 0:
                    payload = {key: value for key, value in result.items() if key not in ("task", "start_worker")}
                    json_response(self, payload, 400 if result["code"] < 5000 else 500)
                    return
                task = result["task"]
                api_code_response(
                    self,
                    result["code"],
                    result["message"],
                    task_id=task["id"],
                    tiktok_blogger_id=task["tiktok_blogger_id"],
                    status=task["status"],
                    min_video_count=task["min_video_count"],
                    available_video_count=task["available_video_count"],
                )
            except Exception as exc:
                api_code_response(self, 5000, str(exc), http_status=500)
            return

        if parsed.path != "/api/run":
            error_response(self, "Not found", 404)
            return
        try:
            body = read_json_body(self)
            prompt1 = body.get("prompt1", "")
            prompt2 = body.get("prompt2", "")
            prompt3 = body.get("prompt3", "")
            prompt4 = body.get("prompt4", "")
            prompt5 = body.get("prompt5", "")
            count = int(body.get("count", 30))
            videos = int(body.get("videos", 15))
            concurrency = max(1, min(int(body.get("concurrency", 200)), 500))
            api_concurrency = max(1, min(int(body.get("apiConcurrency", current_api_concurrency())), 500))
            if not prompt1 or not prompt2 or not prompt3 or not prompt4 or not prompt5:
                raise ValueError("prompt1, prompt2, prompt3, prompt4 and prompt5 are required")
            bloggers = reset_blogger_results(get_or_create_batch(count, videos, force=False))
            job_id = f"job_{int(time.time())}_{random.randint(1000, 9999)}"
            job = {
                "id": job_id,
                "status": "pending",
                "error": "",
                "created_at": time.time(),
                "prompts": {
                    "prompt1": prompt1,
                    "prompt2": prompt2,
                    "prompt3": prompt3,
                    "prompt4": prompt4,
                    "prompt5": prompt5,
                },
                "concurrency": concurrency,
                "api_concurrency": api_concurrency,
                "bloggers": bloggers,
            }
            with JOBS_LOCK:
                JOBS[job_id] = job
            threading.Thread(
                target=run_job,
                args=(job_id, prompt1, prompt2, prompt3, prompt4, prompt5, concurrency, api_concurrency),
                daemon=True,
            ).start()
            json_response(self, job)
        except Exception as exc:
            error_response(self, str(exc), 500)


def main():
    port = int(os.getenv("PORT", "4190"))
    threading.Thread(target=recover_incomplete_tagging_tasks, daemon=True).start()
    start_queue_workers()
    server = ThreadingHTTPServer((os.getenv("HOST", "0.0.0.0"), port), Handler)
    print(f"Serving on http://{os.getenv('HOST', '0.0.0.0')}:{port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
